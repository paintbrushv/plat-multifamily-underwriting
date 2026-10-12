"""The exit-cap sensitivity grid (ADR 0003 §3).

Every run reports the base case at a named set of exit-cap offsets around the
exit cap it ran with (``HOUSE_EXIT_CAP_SENSITIVITY_OFFSETS_BPS``). Each point
pins its exit cap: it is never re-anchored to a price. The exit cap only
changes the reversion, so each point reuses the base cash flows and recomputes
the sale and the returns. Synthetic deals only.
"""

from __future__ import annotations

import copy
import json
from datetime import date
from functools import lru_cache

import pytest
from openpyxl import load_workbook

from engine import underwriting_policy as house
from engine.assumptions import resolve_assumptions
from engine.backsolve import backsolve_price, format_cli_summary
from engine.engine import run_underwriting
from engine.excel_output import write_results_to_workbook
from engine.market_data import MarketDataRegistry, SyntheticProvider
from tests import deal_library as dl

OFFSETS = (-50, -25, 0, 25, 50, 75, 100, 150)
SLUG = "a_stabilized_1990s"
BENCHMARK = {"rate": "0.04", "as_of": "2026-10-03", "source": "synthetic:exit-cap-grid"}
POLICY = {"version": "plat.backsolve-policy/1", "strategy": "cashflow"}


def _grid(results):
    return results["exit_cap_sensitivity"]


def _points(results):
    return _grid(results)["points"]


def _pinned(deal, rate):
    out = copy.deepcopy(deal)
    out["exit_assumptions"]["exit_cap_rate"] = rate
    return out


def test_the_offsets_are_named_house_policy():
    assert house.HOUSE_EXIT_CAP_SENSITIVITY_OFFSETS_BPS == OFFSETS


def test_every_run_reports_the_grid_at_the_named_offsets(minimal_deal_inputs):
    results = run_underwriting(copy.deepcopy(minimal_deal_inputs))
    grid = _grid(results)
    base = minimal_deal_inputs["exit_assumptions"]["exit_cap_rate"]
    assert grid["base_exit_cap_rate"] == base
    assert grid["offsets_bps"] == list(OFFSETS)
    assert [p["offset_bps"] for p in grid["points"]] == list(OFFSETS)
    assert [p["exit_cap_rate"] for p in grid["points"]] == pytest.approx([base + o / 10_000 for o in OFFSETS])
    assert [p["is_base"] for p in grid["points"]] == [o == 0 for o in OFFSETS]
    assert "pinned" in grid["rule"] and "never re-anchored" in grid["rule"]
    for point in grid["points"]:
        assert set(point) >= {
            "offset_bps", "exit_cap_rate", "exit_value", "sale_costs", "net_sale_proceeds",
            "unlevered_irr", "levered_irr", "unlevered_em", "levered_em", "is_base",
        }
    # Year-1 cash-on-cash does not move with the exit cap: shown once.
    assert grid["cash_on_cash_year_1"] == results["metrics"]["coc"]["cash_on_cash_year_1"]
    assert "cash_on_cash_year_1" not in grid["points"][0]


def test_returns_fall_as_the_exit_cap_rises(minimal_deal_inputs):
    points = _points(run_underwriting(copy.deepcopy(minimal_deal_inputs)))
    for field in ("levered_irr", "unlevered_irr", "exit_value", "net_sale_proceeds"):
        values = [p[field] for p in points]
        assert all(a > b for a, b in zip(values, values[1:])), field
    for field in ("levered_em", "unlevered_em"):
        values = [p[field] for p in points]
        assert all(a >= b for a, b in zip(values, values[1:])), field


def test_the_base_point_is_the_base_case(minimal_deal_inputs):
    results = run_underwriting(copy.deepcopy(minimal_deal_inputs))
    base = next(p for p in _points(results) if p["offset_bps"] == 0)
    metrics = results["metrics"]
    assert base["levered_irr"] == metrics["irr"]["levered_irr"]
    assert base["unlevered_irr"] == metrics["irr"]["unlevered_irr"]
    assert base["levered_em"] == metrics["equity_multiple"]["levered_em"]
    assert base["unlevered_em"] == metrics["equity_multiple"]["unlevered_em"]
    assert base["exit_value"] == metrics["exit"]["gross_sale_price"]
    assert base["sale_costs"] == metrics["exit"]["sale_costs"]
    assert base["net_sale_proceeds"] == metrics["exit"]["net_sale_proceeds"]


def test_each_point_matches_a_full_run_with_the_exit_cap_pinned(minimal_deal_inputs):
    results = run_underwriting(copy.deepcopy(minimal_deal_inputs))
    for point in _points(results):
        rerun = run_underwriting(_pinned(minimal_deal_inputs, point["exit_cap_rate"]))["metrics"]
        assert point["levered_irr"] == pytest.approx(rerun["irr"]["levered_irr"], abs=1e-4)
        assert point["unlevered_irr"] == pytest.approx(rerun["irr"]["unlevered_irr"], abs=1e-4)
        assert point["levered_em"] == pytest.approx(rerun["equity_multiple"]["levered_em"], abs=0.01)
        assert point["exit_value"] == pytest.approx(rerun["exit"]["gross_sale_price"], abs=1.0)
        assert point["net_sale_proceeds"] == pytest.approx(rerun["exit"]["net_sale_proceeds"], abs=1.0)


def test_the_grid_can_be_turned_off(minimal_deal_inputs):
    results = run_underwriting(copy.deepcopy(minimal_deal_inputs), exit_cap_sensitivity=False)
    assert "exit_cap_sensitivity" not in results


# ── the anchor: the exit cap as run, pinned ──────────────────────────────


@lru_cache(maxsize=None)
def _library(slug: str, market: bool = False) -> str:
    return json.dumps(dl.ingest(dl.load_deal(slug), market="deal" if market else None))


def test_an_analyst_exit_cap_anchors_the_grid():
    canonical = resolve_assumptions(json.loads(_library(SLUG)), overrides={"exit_cap_rate": 0.0575})
    grid = _grid(run_underwriting(canonical))
    assert grid["base_exit_cap_rate"] == 0.0575
    assert [p["exit_cap_rate"] for p in grid["points"]] == pytest.approx(
        [0.0525, 0.055, 0.0575, 0.06, 0.0625, 0.065, 0.0675, 0.0725])


def test_the_xlsx_has_an_exit_cap_sensitivity_sheet(tmp_path, minimal_deal_inputs):
    results = run_underwriting(copy.deepcopy(minimal_deal_inputs))
    path = tmp_path / "out.xlsx"
    write_results_to_workbook(results, path, minimal_deal_inputs)
    ws = load_workbook(path)["Exit_Cap_Sensitivity"]
    rows = [[c.value for c in row] for row in ws.iter_rows()]
    header = next(i for i, row in enumerate(rows) if row and row[0] == "Offset (bps)")
    body = [row for row in rows[header + 1:] if row and isinstance(row[0], (int, float))]
    assert [row[0] for row in body] == list(OFFSETS)
    labels = rows[header]
    lev = labels.index("Levered IRR")
    assert [row[lev] for row in body] == [p["levered_irr"] for p in _points(results)]


# ── backsolve: the grid at the solved price, never re-anchored ───────────


def _market_registry(cap_rate: float) -> MarketDataRegistry:
    deal = dl.load_deal(SLUG)
    series = {f"2031Q{q}": cap_rate for q in range(1, 5)}
    provider = SyntheticProvider({("market_cap_rate", dl.deal_geo(deal, "submarket")): series},
                                 as_of=date(2026, 6, 30), confidence="high")
    return MarketDataRegistry([provider])


def test_backsolve_reports_the_grid_at_the_solved_price_without_re_anchoring():
    # A low market cap rate makes the solved price a steal: the base exit cap is
    # entry-anchored and compressed at the solved price, and the grid pins
    # offsets around it rather than re-measuring a going-in cap per point.
    canonical = dl.ingest(dl.load_deal(SLUG), market=_market_registry(0.045))
    result = backsolve_price(
        canonical, target_coc="0.06", policy=POLICY, benchmark=BENCHMARK,
        min_price="1000000.00", max_price="60000000.00",
    )
    assert result["status"] == "converged"
    summary, case = result["summary"], result["case"]
    grid = summary["exit_cap_sensitivity"]
    base = case["exit_assumptions"]["exit_cap_rate"]
    assert summary["exit_cap_calculation"]["mode"] == "entry_anchored"
    assert summary["exit_cap_calculation"]["steal"] is True
    assert grid["base_exit_cap_rate"] == base
    assert grid["purchase_price"] == summary["solved_purchase_price"]
    assert [p["exit_cap_rate"] for p in grid["points"]] == pytest.approx([base + o / 10_000 for o in OFFSETS])
    assert grid == result["results"]["exit_cap_sensitivity"]
    base_point = next(p for p in grid["points"] if p["offset_bps"] == 0)
    assert base_point["levered_irr"] == summary["levered_irr"]
    # A point is the solved case with its exit cap pinned, at the solved price.
    point = next(p for p in grid["points"] if p["offset_bps"] == 100)
    rerun = run_underwriting(_pinned(case, point["exit_cap_rate"]), exit_cap_sensitivity=False)
    assert point["levered_irr"] == pytest.approx(rerun["metrics"]["irr"]["levered_irr"], abs=1e-4)
    assert "Exit cap grid" in format_cli_summary(summary, "summary.json")
