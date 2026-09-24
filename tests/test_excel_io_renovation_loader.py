"""Tests for engine.excel_io renovation_programs collision check (Wave 3 / Bug 1.9).

The loader builds `renovation_programs` from `tbl_renovation_programs` but
previously did not check `output_cohort` uniqueness against `unit_cohorts`
loaded earlier from `tbl_unit_cohorts`. Bug 1.9 in the schema audit adds
that check at the loader boundary so the engine never sees a colliding
canonical (mirrors `engine.validator._check_renovation_programs` rule (b)).
"""
import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

from engine.excel_io import load_inputs_from_workbook


def _add_table(ws, name: str, top_left: str, headers, rows):
    start_col = ws[top_left].column
    start_row = ws[top_left].row
    for j, h in enumerate(headers):
        ws.cell(row=start_row, column=start_col + j, value=h)
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row):
            ws.cell(row=start_row + i, column=start_col + j, value=v)
    end_row = start_row + len(rows)
    end_col = start_col + len(headers) - 1
    ref = f"{ws.cell(row=start_row, column=start_col).coordinate}:{ws.cell(row=end_row, column=end_col).coordinate}"
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
    ws.add_table(t)


def _build_minimal_workbook(unit_cohort_ids, reno_programs):
    """Build a minimal valid workbook fixture with the two cohorts named
    `unit_cohort_ids` and the supplied `reno_programs` rows."""
    wb = Workbook()
    ws = wb.active
    ws.title = "INPUTS_CORE"

    _add_table(
        ws,
        "tbl_metadata",
        "A1",
        ["deal_id", "run_id", "as_of_date", "analyst", "purpose", "notes"],
        [["d1", "r1", date(2026, 1, 1), "a", "Screening", ""]],
    )
    _add_table(
        ws,
        "tbl_time_grid",
        "A4",
        ["analysis_start_date", "analysis_end_date"],
        [[date(2026, 1, 1), date(2026, 1, 1)]],
    )
    _add_table(
        ws,
        "tbl_unit_cohorts",
        "A7",
        ["cohort_id", "unit_type", "unit_count", "sqft", "initial_inplace_rent"],
        [[cid, "1B", 10, 650, 1000] for cid in unit_cohort_ids],
    )

    ws2 = wb.create_sheet("INPUTS_RENT")
    _add_table(
        ws2,
        "tbl_market_rent_curve",
        "A1",
        ["cohort_id", "start_period", "end_period", "market_rent"],
        [[unit_cohort_ids[0], "2026-01", "2026-01", 1200]],
    )
    _add_table(
        ws2,
        "tbl_loss_to_lease",
        "A4",
        ["cohort_id", "start_period", "end_period", "ltl_percent"],
        [[unit_cohort_ids[0], "2026-01", "2026-01", 0.1]],
    )
    _add_table(
        ws2,
        "tbl_physical_vacancy_curve",
        "A7",
        ["cohort_id", "start_period", "end_period", "vacancy_rate"],
        [[unit_cohort_ids[0], "2026-01", "2026-01", 0.2]],
    )
    _add_table(
        ws2,
        "tbl_collection_loss_curve",
        "A10",
        ["applies_to", "start_period", "end_period", "loss_rate"],
        [["ALL", "2026-01", "2026-01", 0.05]],
    )

    ws3 = wb.create_sheet("INPUTS_PROGRAMS")
    _add_table(
        ws3,
        "tbl_revenue_programs",
        "A1",
        [
            "program_id",
            "program_name",
            "program_type",
            "pricing_type",
            "price_value",
            "eligible_units",
            "start_period",
            "end_period",
        ],
        [["parking", "Parking", "tenant-based", "$/unit", 50, "ALL", "2026-01", "2026-01"]],
    )
    _add_table(
        ws3,
        "tbl_program_adoption_curve",
        "A4",
        ["program_id", "start_period", "end_period", "adoption_rate"],
        [["parking", "2026-01", "2026-01", 0.5]],
    )

    ws4 = wb.create_sheet("INPUTS_RENO")
    reno_headers = [
        "program_id",
        "program_name",
        "target_cohort",
        "output_cohort",
        "renovation_cost_per_unit",
        "rent_premium_monthly",
        "downtime_days",
        "strategy",
        "start_month",
        "end_month",
        "monthly_pace",
    ]
    _add_table(
        ws4,
        "tbl_renovation_programs",
        "A1",
        reno_headers,
        [
            [
                p["program_id"],
                p["program_name"],
                p["target_cohort"],
                p["output_cohort"],
                p["renovation_cost_per_unit"],
                p["rent_premium_monthly"],
                p["downtime_days"],
                p["strategy"],
                p["start_month"],
                p.get("end_month", ""),
                p["monthly_pace"],
            ]
            for p in reno_programs
        ],
    )

    return wb


class TestExcelIORenovationLoader(unittest.TestCase):
    def test_workbook_renovation_collision_with_unit_cohort_rejected(self):
        """Output_cohort 'A' collides with unit_cohorts cohort_id 'A' → ValueError."""
        unit_cohort_ids = ["A", "B"]
        reno_programs = [
            {
                "program_id": "p1",
                "program_name": "P1",
                "target_cohort": "B",
                # COLLISION: 'A' already exists as a unit_cohort cohort_id.
                "output_cohort": "A",
                "renovation_cost_per_unit": 5000,
                "rent_premium_monthly": 100,
                "downtime_days": 30,
                "strategy": "on_turnover",
                "start_month": "2026-01",
                "end_month": "2026-01",
                "monthly_pace": 1,
            }
        ]
        wb = _build_minimal_workbook(unit_cohort_ids, reno_programs)
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "deal.xlsx"
            wb.save(path)
            with self.assertRaises(ValueError) as cm:
                load_inputs_from_workbook(path)
        msg = str(cm.exception)
        self.assertIn("output_cohort", msg)
        self.assertIn("'A'", msg)

    def test_workbook_renovation_no_collision_loads_clean(self):
        """Non-colliding output_cohort 'A_postreno' loads without error."""
        unit_cohort_ids = ["A", "B"]
        reno_programs = [
            {
                "program_id": "p1",
                "program_name": "P1",
                "target_cohort": "A",
                "output_cohort": "A_postreno",
                "renovation_cost_per_unit": 5000,
                "rent_premium_monthly": 100,
                "downtime_days": 30,
                "strategy": "on_turnover",
                "start_month": "2026-01",
                "end_month": "2026-01",
                "monthly_pace": 1,
            }
        ]
        wb = _build_minimal_workbook(unit_cohort_ids, reno_programs)
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "deal.xlsx"
            wb.save(path)
            inputs = load_inputs_from_workbook(path)

        self.assertIn("renovation_programs", inputs)
        self.assertEqual(len(inputs["renovation_programs"]), 1)
        self.assertEqual(inputs["renovation_programs"][0]["output_cohort"], "A_postreno")


if __name__ == "__main__":
    unittest.main()
