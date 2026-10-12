"""runs/run_deal.py end to end on a synthetic deal-library deal (subprocess)."""

from __future__ import annotations

import copy
import json
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

from tests.deal_library import ingest, load_deal

REPO_ROOT = Path(__file__).resolve().parent.parent
_venv_python = REPO_ROOT / ".venv" / "bin" / "python"
PYTHON = _venv_python if _venv_python.exists() else Path(sys.executable)
SLUG = "b_stabilized_2010s_no_market"


@pytest.fixture(scope="module")
def canonical(tmp_path_factory) -> dict:
    return ingest(load_deal(SLUG))


@pytest.fixture
def canonical_path(canonical, tmp_path) -> Path:
    path = tmp_path / "deal.json"
    path.write_text(json.dumps(canonical), encoding="utf-8")
    return path


def _run(*args: str) -> subprocess.CompletedProcess:
    return subprocess.run(
        [str(PYTHON), "runs/run_deal.py", *args],
        capture_output=True,
        text=True,
        cwd=str(REPO_ROOT),
    )


def test_run_writes_results_sheet_and_workbook(canonical_path, tmp_path):
    out = tmp_path / "run"
    proc = _run(str(canonical_path), "--output-dir", str(out), "--xlsx", "--no-market-data")
    assert proc.returncode == 0, proc.stdout + proc.stderr

    results = json.loads((out / "results.json").read_text(encoding="utf-8"))
    metrics = results["metrics"]
    assert metrics["noi"]["year_1_noi"] > 0
    assert json.loads((out / "validation.json").read_text(encoding="utf-8"))["status"] == "PASS"
    sheet = json.loads((out / "assumptions.json").read_text(encoding="utf-8"))
    assert {row["key"] for row in sheet["rows"]} >= {"rent_growth", "exit_cap_rate"}
    assert json.loads((out / "canonical.json").read_text(encoding="utf-8"))["metadata"]["deal_id"]

    workbook = load_workbook(out / "underwriting.xlsx", read_only=True)
    assert "Assumptions" in workbook.sheetnames
    workbook.close()

    for label in ("NOI (year 1)", "Unlevered IRR", "Levered IRR", "Levered equity multiple",
                  "Minimum DSCR", "Exit value"):
        assert label in proc.stdout
    assert f"{metrics['irr']['levered_irr']:.2%}" in proc.stdout
    assert f"${metrics['exit']['gross_sale_price']:,.0f}" in proc.stdout


def test_run_is_deterministic(canonical_path, tmp_path):
    first, second = tmp_path / "one", tmp_path / "two"
    assert _run(str(canonical_path), "--output-dir", str(first)).returncode == 0
    assert _run(str(canonical_path), "--output-dir", str(second)).returncode == 0
    for name in ("results.json", "assumptions.json", "assumptions.html"):
        assert (first / name).read_bytes() == (second / name).read_bytes()
    assert not (first / "underwriting.xlsx").exists()


def test_assumption_overrides_are_resolved_before_the_run(canonical_path, tmp_path):
    overrides = tmp_path / "overrides.json"
    overrides.write_text(json.dumps({"exit_cap_rate": {"value": 0.0625, "note": "synthetic test"}}), encoding="utf-8")
    out = tmp_path / "run"
    proc = _run(str(canonical_path), "--output-dir", str(out), "--no-market-data",
                "--assumption-overrides", str(overrides))
    assert proc.returncode == 0, proc.stdout + proc.stderr
    assert "Applied 1 assumption override(s): exit_cap_rate" in proc.stdout

    results = json.loads((out / "results.json").read_text(encoding="utf-8"))
    assert results["metrics"]["yields"]["exit_cap_rate"] == 0.0625
    row = next(r for r in json.loads((out / "assumptions.json").read_text(encoding="utf-8"))["rows"]
               if r["key"] == "exit_cap_rate")
    assert row["value"] == 0.0625
    assert row["source"] == "analyst"
    assert row["override_note"] == "synthetic test"


def test_validation_failure_exits_1_with_the_report(canonical, tmp_path):
    broken = copy.deepcopy(canonical)
    broken["metadata"]["property_summary"].pop("property_tax_policy")
    path = tmp_path / "no_millage.json"
    path.write_text(json.dumps(broken), encoding="utf-8")
    out = tmp_path / "run"
    proc = _run(str(path), "--output-dir", str(out))
    assert proc.returncode == 1, proc.stdout + proc.stderr
    assert "ERROR [PROPERTY_TAX_MILLAGE_MISSING]" in proc.stdout
    assert "validation failed" in proc.stderr
    report = json.loads((out / "validation.json").read_text(encoding="utf-8"))
    assert report["status"] == "FAIL"
    assert not (out / "results.json").exists()


def test_engine_refusal_exits_1(canonical, tmp_path):
    placeholder = copy.deepcopy(canonical)
    placeholder["exit_assumptions"]["exit_cap_rate"] = None
    path = tmp_path / "placeholder.json"
    path.write_text(json.dumps(placeholder), encoding="utf-8")
    proc = _run(str(path), "--output-dir", str(tmp_path / "run"))
    assert proc.returncode == 1, proc.stdout + proc.stderr
    assert "EXIT_CAP_RATE_UNRESOLVED" in proc.stderr
    assert not (tmp_path / "run" / "results.json").exists()


def test_bad_inputs_exit_2(canonical_path, tmp_path):
    missing = _run(str(tmp_path / "missing.json"), "--output-dir", str(tmp_path / "a"))
    assert missing.returncode == 2
    assert "canonical JSON not found" in missing.stderr

    garbage = tmp_path / "garbage.json"
    garbage.write_text("{not json", encoding="utf-8")
    assert _run(str(garbage), "--output-dir", str(tmp_path / "b")).returncode == 2

    not_object = tmp_path / "list.json"
    not_object.write_text("[]", encoding="utf-8")
    assert _run(str(not_object), "--output-dir", str(tmp_path / "c")).returncode == 2

    unknown = tmp_path / "unknown_override.json"
    unknown.write_text(json.dumps({"not_an_assumption": 1}), encoding="utf-8")
    proc = _run(str(canonical_path), "--output-dir", str(tmp_path / "d"), "--no-market-data",
                "--assumption-overrides", str(unknown))
    assert proc.returncode == 2
    assert "failed to apply assumption overrides" in proc.stderr

    no_output_dir = _run(str(canonical_path))
    assert no_output_dir.returncode == 2
