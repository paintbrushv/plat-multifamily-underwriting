"""Tests for engine.rediq_bridge.extract_renovation_programs naming convention.

Wave 3 / Naming alignment: federation-emitted output_cohorts must use the
``_postreno`` suffix (matching plat-agent's `schema_mapper`) to avoid
colliding with rent-roll-derived ``*_renovated`` subtotal cohort_ids
(the the schema audit Family A bug).

Prior behavior emitted ``f"{fp_name}_reno"`` which is also ambiguous when a
canonical contains both rent-roll-derived ``Beal Renovated`` cohorts and
federation-spliced renovation programs targeting ``Beal``.
"""
import unittest

from openpyxl import Workbook

from engine.rediq_bridge import extract_renovation_programs


def _build_renovation_workbook(fp_name="1BR-A", units_yr1=12):
    """Construct a minimal RedIQ-shaped workbook with one renovation row.

    Layout matches REDIQ_LAYOUT['renovations']:
      - sheet "Input"
      - renovations_start_row = 87 (default)
      - row 87: B=fp_name, C=downtime, D=cost, E=premium, F=post_market,
                G..M = yearly unit counts (Yr1..Yr7)
      - inflation row 169, cols G..Q hold annual inflation rates (we leave
        blank → reno_inflation defaults to 0).
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Input"

    reno_row = 87
    ws[f"B{reno_row}"] = fp_name
    ws[f"C{reno_row}"] = 30  # downtime days
    ws[f"D{reno_row}"] = 5000  # cost per unit
    ws[f"E{reno_row}"] = 100  # rent premium
    ws[f"F{reno_row}"] = 1300  # post-renovation market rent
    # Yr1 unit count → column G (col 7)
    ws.cell(row=reno_row, column=7, value=units_yr1)

    return wb


class TestRediqBridgePostrenoNaming(unittest.TestCase):
    def test_extract_renovations_uses_postreno_suffix(self):
        """Federation-emitted output_cohort must end with '_postreno' (not
        '_reno' and not '_renovated')."""
        wb = _build_renovation_workbook(fp_name="Beal", units_yr1=10)
        programs = extract_renovation_programs(wb, start_date="2026-01-01")

        self.assertGreaterEqual(len(programs), 1, "Expected at least one renovation program")
        for p in programs:
            output = p.get("output_cohort")
            self.assertIsNotNone(output, "output_cohort missing")
            self.assertTrue(
                output.endswith("_postreno"),
                f"output_cohort '{output}' does not end with '_postreno' "
                "(see the schema audit / Wave 2 schema_mapper convention)",
            )
            # Negative assertions for the historic suffixes:
            self.assertFalse(
                output.endswith("_reno") and not output.endswith("_postreno"),
                f"output_cohort '{output}' still uses legacy '_reno' suffix",
            )
            self.assertFalse(
                output.endswith("_renovated"),
                f"output_cohort '{output}' uses '_renovated' suffix "
                "(reserved for rent-roll-derived subtotal cohort_ids)",
            )
            # Suffix is built from the floor-plan name.
            self.assertEqual(output, "Beal_postreno")

    def test_target_cohort_unchanged(self):
        """target_cohort should remain the raw floor-plan name (no suffix)."""
        wb = _build_renovation_workbook(fp_name="Beal", units_yr1=10)
        programs = extract_renovation_programs(wb, start_date="2026-01-01")
        self.assertGreaterEqual(len(programs), 1)
        for p in programs:
            self.assertEqual(p["target_cohort"], "Beal")


if __name__ == "__main__":
    unittest.main()
