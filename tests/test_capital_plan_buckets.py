"""ADR 0003 §3: the three capex buckets on every deal.

1. Immediate / deferred capex at close (``capital_plan.immediate_capex``):
   funded with equity at close (or a lender holdback), in the Day-0 basis,
   never in NOI or the monthly cash flows, separate from the renovation budget.
2. Recurring capex by vintage (``capital_plan.recurring_capex``): $/unit below
   NOI, grown every analysis year. PCA > analyst > house vintage band (low
   end stabilized, midpoint value-add, high end with deferred maintenance).
3. The lender replacement reserve ($250/unit) is a floor, not an addition:
   the charge below NOI each month is max(recurring capex, reserve).

All deals and figures are synthetic.
"""
from __future__ import annotations

import copy
import json
import sys
from functools import lru_cache

import pytest

from engine import underwriting_policy as house
from engine.assumptions import resolve_assumptions
from engine.assumptions import guardrails
from engine.assumptions.guardrails import GuardrailProfile
from engine.assumptions.rationale import rationale_for
from engine.assumptions.register import build_register
from engine.assumptions.register_html import render_register_html
from engine.assumptions.sheet import (
    read_assumption_overrides,
    read_assumption_overrides_xlsx,
    sheet_rows,
    write_assumptions_json,
    write_assumptions_xlsx,
)
from engine.backsolve import backsolve_price
from engine.engine import run_underwriting
from engine.validator import validate_deal
from tests import deal_library as dl

UNITS = 100  # the minimal deal fixture
BENCHMARK = {"rate": "0.04", "as_of": "2026-10-01", "source": "synthetic:capex-buckets"}
STABILIZED = "a_stabilized_1990s"  # 200 units, 1990s band $750-$2,000


def _with_plan(deal, **plan):
    out = copy.deepcopy(deal)
    out["capital_plan"] = plan
    return out


def _by_year(results, field):
    return [row[field] for row in results["capital_plan"]["below_noi_capex_charge"]["by_analysis_year"]]


# ── engine: where each bucket enters ─────────────────────────────────────


def test_without_a_capital_plan_the_reserve_is_charged_exactly_as_before(minimal_deal_inputs):
    results = run_underwriting(copy.deepcopy(minimal_deal_inputs))
    assert "capital_plan" not in results
    months = results["cashflow"]["by_month"]
    assert {m["replacement_reserves"] for m in months} == {round(25_000 / 12, 2)}
    assert "recurring_capex_by_vintage" not in months[0]


@pytest.mark.parametrize("recurring, binding", [(200.0, "replacement_reserve"), (500.0, "recurring_capex")])
def test_the_reserve_is_a_floor_not_stacked(minimal_deal_inputs, recurring, binding):
    base = run_underwriting(copy.deepcopy(minimal_deal_inputs))
    results = run_underwriting(_with_plan(
        minimal_deal_inputs, recurring_capex={"per_unit_annual": recurring, "growth_rate": 0.0}))
    charge = max(recurring, 250.0) * UNITS
    years = results["capital_plan"]["below_noi_capex_charge"]["by_analysis_year"]
    assert [y["below_noi_capex_charge"] for y in years] == pytest.approx([charge] * 5, abs=0.1)
    assert [y["recurring_capex_by_vintage"] for y in years] == pytest.approx([recurring * UNITS] * 5, abs=0.1)
    assert [y["replacement_reserve_floor"] for y in years] == pytest.approx([25_000.0] * 5, abs=0.1)
    assert {y["binding"] for y in years} == {binding}
    # Charged once: never recurring + reserve.
    month = results["cashflow"]["by_month"][0]
    assert month["replacement_reserves"] == pytest.approx(charge / 12, abs=0.01)
    assert month["replacement_reserves"] < month["recurring_capex_by_vintage"] + month["replacement_reserve_floor"]
    # NOI and DSCR are before capex; unlevered cash flow carries the charge.
    for field in ("net_operating_income", "debt_service"):
        assert [y[field] for y in results["cashflow"]["by_year"]] == [y[field] for y in base["cashflow"]["by_year"]]
    assert results["metrics"]["dscr"] == base["metrics"]["dscr"]
    assert results["metrics"]["yields"]["going_in_cap_rate"] == base["metrics"]["yields"]["going_in_cap_rate"]
    extra = charge - 25_000.0
    assert results["metrics"]["coc"]["components"]["replacement_reserves_year_1"] == pytest.approx(
        base["metrics"]["coc"]["components"]["replacement_reserves_year_1"] + extra, abs=1.0)
    if extra:
        assert results["metrics"]["coc"]["cash_on_cash_year_1"] < base["metrics"]["coc"]["cash_on_cash_year_1"]
        assert results["metrics"]["irr"]["levered_irr"] < base["metrics"]["irr"]["levered_irr"]
    assert results["capital_plan"]["metric_basis"]["dscr"].startswith("NOI / debt service (before capex)")


def test_recurring_capex_grows_every_year_and_can_overtake_the_reserve(minimal_deal_inputs):
    results = run_underwriting(_with_plan(
        minimal_deal_inputs, recurring_capex={"per_unit_annual": 240.0, "growth_rate": 0.10}))
    assert _by_year(results, "recurring_capex_by_vintage") == pytest.approx(
        [240 * UNITS * 1.1 ** k for k in range(5)], abs=0.5)
    assert _by_year(results, "binding") == ["replacement_reserve"] + ["recurring_capex"] * 4
    assert _by_year(results, "below_noi_capex_charge")[0] == pytest.approx(25_000.0, abs=0.1)


def test_immediate_capex_is_equity_at_close_never_noi(minimal_deal_inputs):
    base = run_underwriting(copy.deepcopy(minimal_deal_inputs))
    results = run_underwriting(_with_plan(minimal_deal_inputs, immediate_capex={"amount": 400_000.0}))
    # The fixture states its total equity basis; the bucket is added to it, never assumed inside it.
    stated = minimal_deal_inputs["purchase_assumptions"]["total_equity_basis"]
    assert results["metrics"]["coc"]["components"]["equity_basis"] == pytest.approx(stated + 400_000.0)
    assert results["capital_plan"]["immediate_capex"] == {
        "amount": 400_000.0, "funding": "equity_at_close", "in_total_equity_basis": True,
        "in_total_unlevered_basis": True, "amount_in_stated_basis": 0.0, "added_by_engine": 400_000.0,
    }
    # Not in NOI, the below-NOI charge or operating capex: only the Day-0 outlay moves.
    for field in ("net_operating_income", "replacement_reserves", "total_capex", "leveraged_cash_flow"):
        assert [y[field] for y in results["cashflow"]["by_year"]] == [y[field] for y in base["cashflow"]["by_year"]]
    assert results["metrics"]["coc"]["cash_on_cash_year_1"] < base["metrics"]["coc"]["cash_on_cash_year_1"]
    assert results["metrics"]["irr"]["levered_irr"] < base["metrics"]["irr"]["levered_irr"]
    assert results["metrics"]["irr"]["unlevered_irr"] < base["metrics"]["irr"]["unlevered_irr"]

    # Without a stated basis the default (equity contribution) carries it.
    unstated = _with_plan(minimal_deal_inputs, immediate_capex={"amount": 400_000.0})
    unstated["purchase_assumptions"].pop("total_equity_basis")
    equity = minimal_deal_inputs["purchase_assumptions"]["equity_contribution"]
    assert run_underwriting(unstated)["metrics"]["coc"]["components"]["equity_basis"] == pytest.approx(
        equity + 400_000.0)
    # A basis that already carries part of it (backsolve's priced case) gets only the rest.
    partial = _with_plan(minimal_deal_inputs,
                         immediate_capex={"amount": 400_000.0, "amount_in_stated_basis": 150_000.0})
    partial_results = run_underwriting(partial)
    assert partial_results["metrics"]["coc"]["components"]["equity_basis"] == pytest.approx(stated + 250_000.0)
    assert partial_results["capital_plan"]["immediate_capex"]["added_by_engine"] == 250_000.0


def test_a_lender_holdback_moves_the_unlevered_basis_only(minimal_deal_inputs):
    base = run_underwriting(copy.deepcopy(minimal_deal_inputs))
    results = run_underwriting(_with_plan(
        minimal_deal_inputs, immediate_capex={"amount": 400_000.0, "funding": "lender_holdback"}))
    assert results["metrics"]["coc"]["components"]["equity_basis"] == base["metrics"]["coc"]["components"]["equity_basis"]
    assert results["metrics"]["irr"]["levered_irr"] == base["metrics"]["irr"]["levered_irr"]
    assert results["metrics"]["irr"]["unlevered_irr"] < base["metrics"]["irr"]["unlevered_irr"]
    assert results["capital_plan"]["immediate_capex"]["in_total_equity_basis"] is False


def test_immediate_capex_and_the_renovation_budget_are_separate_budgets(minimal_deal_inputs):
    deal = _with_plan(minimal_deal_inputs, immediate_capex={"amount": 100_000.0})
    deal["renovation_capex_funding"] = "equity_at_close"
    deal["renovation_programs"] = [{
        "program_id": "synthetic_interiors", "program_name": "Synthetic interiors",
        "target_cohort": "1BR", "output_cohort": "1BR_reno", "renovation_cost_per_unit": 8_000,
        "rent_premium_monthly": 100, "downtime_days": 30, "strategy": "proactive",
        "start_month": "2026-09", "monthly_pace": 2, "max_units": 10,
    }]
    results = run_underwriting(deal)
    stated = minimal_deal_inputs["purchase_assumptions"]["total_equity_basis"]
    plan = results["capital_plan"]
    assert plan["renovation_budget"] == {"amount": 80_000.0, "funding": "equity_at_close"}
    assert plan["immediate_capex"]["amount"] == 100_000.0
    # The stated basis is taken as carrying the renovation budget (its convention);
    # the immediate-capex bucket is added on top, so the two never double count.
    assert results["metrics"]["coc"]["components"]["equity_basis"] == pytest.approx(stated + 100_000)
    unstated = copy.deepcopy(deal)
    unstated["purchase_assumptions"].pop("total_equity_basis")
    equity = minimal_deal_inputs["purchase_assumptions"]["equity_contribution"]
    assert run_underwriting(unstated)["metrics"]["coc"]["components"]["equity_basis"] == pytest.approx(
        equity + 80_000 + 100_000)


def test_the_workbook_shows_the_buckets_and_the_charge(tmp_path, minimal_deal_inputs):
    from openpyxl import load_workbook

    from engine.excel_output import write_results_to_workbook

    deal = _with_plan(minimal_deal_inputs, immediate_capex={"amount": 400_000.0},
                      recurring_capex={"per_unit_annual": 500.0, "growth_rate": 0.03})
    path = tmp_path / "results.xlsx"
    write_results_to_workbook(run_underwriting(copy.deepcopy(deal)), path, inputs=deal)
    workbook = load_workbook(path)
    labels = {row[0]: row for row in workbook["Summary"].iter_rows(values_only=True) if row[0]}
    assert "CAPITAL PLAN (3 CAPEX BUCKETS)" in labels
    assert labels["1. Immediate capex at close:"][1] == 400_000.0
    assert labels["2. Recurring capex, $/unit yr 1:"][1] == 500.0
    assert labels["3. Replacement reserve, year 1:"][1] == pytest.approx(25_000.0, abs=0.1)
    assert labels["Below-NOI charge, year 1:"][1] == pytest.approx(50_000.0, abs=0.1)
    annual = {row[0] for row in workbook["Annual_ProForma"].iter_rows(values_only=True) if row[0]}
    assert {"Below-NOI CapEx Charge", "  Recurring CapEx (vintage)", "  Replacement Reserve (floor)"} <= annual
