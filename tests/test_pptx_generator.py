"""
Tests for PowerPoint Presentation Generator
"""
import tempfile
import unittest
from pathlib import Path

from engine.pptx_generator import generate_presentation


def _make_inputs():
    return {
        "metadata": {"deal_id": "Test Property", "city": "Dallas", "state": "TX"},
        "time_grid": {"analysis_start_date": "2026-01-01", "analysis_end_date": "2030-12-31"},
        "unit_cohorts": [
            {"cohort_id": "1BR", "unit_count": 100, "sqft": 750, "initial_inplace_rent": 1100},
            {"cohort_id": "2BR", "unit_count": 50, "sqft": 1000, "initial_inplace_rent": 1400},
        ],
        "purchase_assumptions": {"purchase_price": 15000000},
        "exit_assumptions": {"exit_cap_rate": 0.055},
        "debt_terms": {"commitment": 10500000, "rate": 0.065, "amort_years": 30, "io_months": 24},
        "growth_assumptions": {"annual_growth_rate": 0.03},
        "physical_vacancy_curve": [{"cohort_id": "1BR", "vacancy_rate": 0.05}],
    }


def _make_results():
    return {
        "metrics": {
            "irr": {"levered_irr": 0.18, "unlevered_irr": 0.12, "partnership_irr": 0.14},
            "equity_multiple": {"levered_em": 2.1, "unlevered_em": 1.75, "partnership_em": 1.65},
            "dscr": {"average_dscr": 1.35, "minimum_dscr": 1.20},
            "yields": {"going_in_cap_rate": 0.05},
            "cash_on_cash": {"by_year": [{"year": "Y1", "yield": 0.06}]},
        },
        "cashflow": {
            "by_year": [
                {"year": "2026", "net_operating_income": 750000, "levered_cashflow": 200000},
                {"year": "2027", "net_operating_income": 785000, "levered_cashflow": 235000},
                {"year": "2028", "net_operating_income": 820000, "levered_cashflow": 270000},
                {"year": "2029", "net_operating_income": 855000, "levered_cashflow": 305000},
                {"year": "2030", "net_operating_income": 890000, "levered_cashflow": 5500000},
            ],
        },
    }


def _make_scenario_results():
    return {
        "comparison": {
            "bull": {
                "metrics": {
                    "levered_irr": 0.22, "unlevered_irr": 0.15,
                    "levered_em": 2.4, "unlevered_em": 1.9,
                    "noi_year_1": 780000, "noi_exit_year": 950000,
                    "average_dscr": 1.40, "going_in_cap": 0.05,
                },
                "assumptions": {
                    "rent_growth": 0.035, "exit_cap_rate": 0.0525, "avg_vacancy": 0.04,
                },
            },
            "base": {
                "metrics": {
                    "levered_irr": 0.18, "unlevered_irr": 0.12,
                    "levered_em": 2.1, "unlevered_em": 1.75,
                    "noi_year_1": 750000, "noi_exit_year": 890000,
                    "average_dscr": 1.35, "going_in_cap": 0.05,
                },
                "assumptions": {
                    "rent_growth": 0.03, "exit_cap_rate": 0.055, "avg_vacancy": 0.05,
                },
            },
            "bear": {
                "metrics": {
                    "levered_irr": 0.13, "unlevered_irr": 0.09,
                    "levered_em": 1.7, "unlevered_em": 1.55,
                    "noi_year_1": 720000, "noi_exit_year": 830000,
                    "average_dscr": 1.25, "going_in_cap": 0.05,
                },
                "assumptions": {
                    "rent_growth": 0.025, "exit_cap_rate": 0.0575, "avg_vacancy": 0.07,
                },
            },
        },
    }


class TestGeneratePresentation(unittest.TestCase):
    def test_creates_pptx_file(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "test.pptx"
            result = generate_presentation(_make_results(), _make_inputs(), path)
            self.assertTrue(result.exists())
            self.assertGreater(result.stat().st_size, 10000)

    def test_has_five_slides(self):
        from pptx import Presentation as PptxReader
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "test.pptx"
            generate_presentation(_make_results(), _make_inputs(), path)
            prs = PptxReader(str(path))
            self.assertEqual(len(prs.slides), 5)

    def test_with_scenarios(self):
        from pptx import Presentation as PptxReader
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "test_scenarios.pptx"
            generate_presentation(
                _make_results(), _make_inputs(), path,
                scenario_results=_make_scenario_results()
            )
            prs = PptxReader(str(path))
            self.assertEqual(len(prs.slides), 5)
            self.assertGreater(path.stat().st_size, 10000)

    def test_widescreen_dimensions(self):
        from pptx import Presentation as PptxReader
        from pptx.util import Inches
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "test.pptx"
            generate_presentation(_make_results(), _make_inputs(), path)
            prs = PptxReader(str(path))
            self.assertEqual(prs.slide_width, Inches(13.333))
            self.assertEqual(prs.slide_height, Inches(7.5))

    def test_cover_slide_has_deal_name(self):
        from pptx import Presentation as PptxReader
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "test.pptx"
            generate_presentation(_make_results(), _make_inputs(), path)
            prs = PptxReader(str(path))
            cover = prs.slides[0]
            texts = [shape.text for shape in cover.shapes if shape.has_text_frame]
            self.assertTrue(any("TEST PROPERTY" in t for t in texts))

    def test_summary_slide_shows_year_built(self):
        from pptx import Presentation as PptxReader
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "test.pptx"
            inputs = _make_inputs()
            inputs["metadata"]["address"] = {
                "street": "123 Main St",
                "city": "Dallas",
                "state": "TX",
                "zip": "75201",
            }
            inputs["metadata"]["year_built"] = 1987
            generate_presentation(_make_results(), inputs, path)
            prs = PptxReader(str(path))
            summary = prs.slides[1]
            texts = []
            for shape in summary.shapes:
                if hasattr(shape, "has_table") and shape.has_table:
                    for row in shape.table.rows:
                        for cell in row.cells:
                            texts.append(cell.text)
                elif shape.has_text_frame:
                    texts.append(shape.text)
            self.assertTrue(any("Year Built" in t for t in texts))
            self.assertTrue(any("1,987" in t for t in texts))
            self.assertTrue(any("Location" in t for t in texts))
            self.assertTrue(any("Dallas, TX" in t for t in texts))

    def test_empty_cashflow_handled(self):
        results = _make_results()
        results["cashflow"]["by_year"] = []
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "test.pptx"
            generate_presentation(results, _make_inputs(), path)
            self.assertTrue(path.exists())

    def test_minimal_inputs(self):
        """Test with minimal inputs (missing optional fields)."""
        inputs = {
            "metadata": {"deal_id": "Minimal Deal"},
            "unit_cohorts": [{"cohort_id": "1BR", "unit_count": 10, "sqft": 700}],
            "purchase_assumptions": {"purchase_price": 1000000},
            "exit_assumptions": {},
            "debt_terms": {},
        }
        results = {
            "metrics": {"irr": {}, "equity_multiple": {}, "dscr": {}, "yields": {}},
            "cashflow": {"by_year": []},
        }
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "test.pptx"
            generate_presentation(results, inputs, path)
            self.assertTrue(path.exists())


if __name__ == "__main__":
    unittest.main()
