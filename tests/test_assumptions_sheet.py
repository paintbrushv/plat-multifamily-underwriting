"""ADR 0003 slice 7: the assumptions sheet (JSON and Excel) with overrides.

Every run writes an assumptions sheet; an analyst edits it and the resolver
re-applies the edits as overrides, keeping each entry's ``default_value``.
"""
from __future__ import annotations

import copy
import json
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

from engine.assumptions import resolve_assumptions
from engine.assumptions.sheet import (
    DEFAULT_EDIT_NOTE,
    ROW_FIELDS,
    SHEET_TITLE,
    SHEET_VERSION,
    TABLE_NAME,
    add_assumptions_sheet,
    assumptions_document,
    delta,
    load_assumption_overrides,
    read_assumption_overrides,
    read_assumption_overrides_xlsx,
    read_assumption_rows_xlsx,
    sheet_path_for,
    sheet_rows,
    write_assumptions_json,
    write_assumptions_xlsx,
)
from engine.ingest.ingestion_pipeline import build_deal_from_documents

REPO_ROOT = Path(__file__).resolve().parent.parent
FIXTURES = REPO_ROOT / "tests" / "fixtures"
ROLL = FIXTURES / "sample_rent_roll.csv"
T12 = FIXTURES / "entrata_style_t12.csv"
_venv_python = REPO_ROOT / ".venv" / "bin" / "python"
VENV_PYTHON = _venv_python if _venv_python.exists() else Path(sys.executable)


def _entry(value, default, **extra):
    entry = {
        "value": value,
        "unit": "decimal_yoy",
        "source": "house:long_run",
        "basis": "synthetic",
        "as_of": None,
        "confidence": "medium",
        "candidates": [{"source": "house:long_run", "value": default}],
        "default_value": default,
        "override": None,
        "review_required": True,
    }
    entry.update(extra)
    return entry


@pytest.fixture()
def synthetic_canonical():
    """Entries of every shape, including keys and fields the sheet has never seen."""
    return {
        "metadata": {"deal_id": "Synthetic Sheet"},
        "assumptions": {
            "resolver_version": "plat.assumptions/1",
            "mode": "underwrite",
            "entries": {
                "zz_custom_mix": _entry({"1br": 0.5, "2br": 0.5}, {"1br": 0.5, "2br": 0.5},
                                        unit="share_by_unit_type"),
                "opex.payroll.base": _entry(
                    1800.0, 1650.0, unit="usd_per_unit_year", per_unit_annual=1800.0, source="analyst",
                    override={"value": 1800.0, "note": "new payroll plan"},
                ),
                "collection_loss": _entry(
                    0.01, 0.005, unit="decimal_of_rent", source="analyst",
                    override={"value": 0.01, "note": "T12 bad debt"},
                    guardrail={"status": "below_floor", "floor": 0.0075, "mode": "underwrite"},
                ),
                "physical_vacancy": _entry(
                    [0.08, 0.06, 0.05], [0.07, 0.06, 0.05], unit="decimal_of_units_by_analysis_year",
                    source="analyst", override={"value": [0.08, 0.06, 0.05], "note": "slow lease-up"},
                ),
                "rent_growth": _entry([0.021, 0.028, 0.03], [0.021, 0.028, 0.03],
                                      unit="decimal_yoy_by_analysis_year", guardrail="ok"),
            },
        },
    }


def _ingest(**kwargs):
    params = dict(
        property_id="Synthetic Sheet",
        rent_roll_path=ROLL,
        t12_path=T12,
        analysis_start="2026-08",
        analysis_end="2031-07",
        year_built=1984,
        purchase_price=25_000_000,
        ltv=0.75,
        loan_rate=0.065,
        amort_years=30,
        exit_cap_rate=0.055,
        sponsor_equity_pct=0.10,
        pref_return=0.08,
    )
    params.update(kwargs)
    return build_deal_from_documents(**params)


@pytest.fixture(scope="module")
def resolved_deal():
    return _ingest()


def _norm(value):
    """Rows compared after an xlsx round trip: xlsx keeps 16 significant digits."""
    if isinstance(value, float):
        return round(value, 12) + 0.0
    if isinstance(value, int) and not isinstance(value, bool):
        return float(value)
    if isinstance(value, list):
        return [_norm(v) for v in value]
    if isinstance(value, dict):
        return {k: _norm(v) for k, v in value.items()}
    return value


def _existing_overrides(deal):
    return {
        key: {"value": entry["override"]["value"], "note": entry["override"]["note"]}
        for key, entry in deal["assumptions"]["entries"].items()
        if entry.get("override")
    }


# ── rows ─────────────────────────────────────────────────────────────────


def test_rows_cover_scalar_vector_dict_entries_and_unknown_fields(synthetic_canonical):
    rows = sheet_rows(synthetic_canonical)
    by_key = {row["key"]: row for row in rows}

    assert [row["key"] for row in rows] == [
        "rent_growth", "collection_loss", "physical_vacancy", "opex.payroll.base", "zz_custom_mix",
    ]
    for row in rows:
        assert tuple(row) == ROW_FIELDS

    vector = by_key["rent_growth"]
    assert vector["value"] == [0.021, 0.028, 0.03]
    assert vector["default_value"] == [0.021, 0.028, 0.03]
    assert vector["delta"] == [0.0, 0.0, 0.0]
    assert vector["guardrail_status"] == "ok"
    assert vector["extra"] == {"guardrail": "ok"}
    assert vector["override_value"] is None and vector["override_note"] is None

    overridden_vector = by_key["physical_vacancy"]
    assert overridden_vector["delta"] == pytest.approx([0.01, 0.0, 0.0])
    assert overridden_vector["override_value"] == [0.08, 0.06, 0.05]
    assert overridden_vector["override_note"] == "slow lease-up"

    scalar = by_key["collection_loss"]
    assert scalar["delta"] == pytest.approx(0.005)
    assert scalar["override_value"] == 0.01
    assert scalar["override_note"] == "T12 bad debt"
    assert scalar["guardrail_status"] == "below_floor"
    assert scalar["extra"] == {"guardrail": {"status": "below_floor", "floor": 0.0075, "mode": "underwrite"}}
    assert scalar["source"] == "analyst" and scalar["review_required"] is True

    unknown_field = by_key["opex.payroll.base"]
    assert unknown_field["extra"] == {"per_unit_annual": 1800.0}
    assert unknown_field["delta"] == 150.0
    assert unknown_field["guardrail_status"] is None

    mapping = by_key["zz_custom_mix"]
    assert mapping["value"] == {"1br": 0.5, "2br": 0.5}
    assert mapping["delta"] is None
    assert mapping["extra"] is None


def test_rows_from_the_resolver_cover_every_entry(resolved_deal):
    rows = sheet_rows(resolved_deal)
    entries = resolved_deal["assumptions"]["entries"]
    assert {row["key"] for row in rows} == set(entries)
    for row in rows:
        entry = entries[row["key"]]
        assert row["value"] == entry["value"]
        assert row["default_value"] == entry["default_value"]
        assert row["source"] == entry["source"]
        assert row["confidence"] == entry["confidence"]
    exit_cap = next(row for row in rows if row["key"] == "exit_cap_rate")
    assert exit_cap["override_value"] == 0.055


def test_a_canonical_without_assumptions_has_no_rows(minimal_deal_inputs):
    assert sheet_rows(minimal_deal_inputs) == []
    assert sheet_rows({"assumptions": {"entries": None}}) == []


@pytest.mark.parametrize(
    "value, default, expected",
    [
        (0.03, 0.025, 0.005),
        (0.025, 0.03, -0.005),
        (0.03, 0.03, 0.0),
        (25_200, 20_000.0, 5_200.0),
        ([0.03, 0.02], [0.025, 0.02], [0.005, 0.0]),
        ((1, 2), [0, 0], [1.0, 2.0]),
        ([0.03, 0.02], [0.025], None),
        ([0.03, None], [0.025, 0.02], None),
        ([0.03, "x"], [0.025, 0.02], None),
        ({"a": 1.0}, {"a": 0.5}, None),
        (None, 0.03, None),
        (0.03, None, None),
        ("text", "other", None),
        (True, False, None),
        ([], [], None),
    ],
)
def test_delta(value, default, expected):
    result = delta(value, default)
    if expected is None:
        assert result is None
    else:
        assert result == pytest.approx(expected)


def test_delta_is_free_of_float_noise_and_negative_zero():
    assert delta(0.03, 0.025) == 0.005
    assert repr(delta(-0.0, 0.0)) == "0.0"
    assert delta([0.1 + 0.2], [0.3]) == [0.0]


# ── JSON ─────────────────────────────────────────────────────────────────


def test_json_round_trip_preserves_vectors_and_dicts(synthetic_canonical, tmp_path):
    path = write_assumptions_json(synthetic_canonical, tmp_path / "assumptions.json")
    document = json.loads(path.read_text(encoding="utf-8"))

    assert document["sheet_version"] == SHEET_VERSION
    assert document["resolver_version"] == "plat.assumptions/1"
    assert document["deal_id"] == "Synthetic Sheet"
    assert document["columns"] == list(ROW_FIELDS)
    assert document["rows"] == sheet_rows(synthetic_canonical)

    # Unedited, exactly the recorded overrides come back.
    recorded = {
        "collection_loss": {"value": 0.01, "note": "T12 bad debt"},
        "physical_vacancy": {"value": [0.08, 0.06, 0.05], "note": "slow lease-up"},
        "opex.payroll.base": {"value": 1800.0, "note": "new payroll plan"},
    }
    assert read_assumption_overrides(path) == recorded

    by_key = {row["key"]: row for row in document["rows"]}
    by_key["rent_growth"]["value"] = [0.02, 0.025, 0.025]
    by_key["zz_custom_mix"]["value"] = {"1br": 0.45, "2br": 0.55}
    by_key["physical_vacancy"]["override_value"] = [0.09, 0.07, 0.05]
    by_key["collection_loss"]["override_value"] = None  # cleared, and the value is the default again
    by_key["collection_loss"]["value"] = 0.005
    path.write_text(json.dumps(document), encoding="utf-8")

    assert read_assumption_overrides(path) == {
        "rent_growth": {"value": [0.02, 0.025, 0.025], "note": DEFAULT_EDIT_NOTE},
        "physical_vacancy": {"value": [0.09, 0.07, 0.05], "note": "slow lease-up"},
        "opex.payroll.base": {"value": 1800.0, "note": "new payroll plan"},
        "zz_custom_mix": {"value": {"1br": 0.45, "2br": 0.55}, "note": DEFAULT_EDIT_NOTE},
    }


def test_override_column_wins_over_an_edited_value(tmp_path):
    rows = [{"key": "rent_growth", "value": 0.04, "default_value": 0.03,
             "override_value": 0.035, "override_note": " IC case "}]
    path = tmp_path / "rows.json"
    path.write_text(json.dumps(rows), encoding="utf-8")
    assert read_assumption_overrides(path) == {"rent_growth": {"value": 0.035, "note": "IC case"}}


def test_override_cohorts_round_trip(tmp_path, resolved_deal):
    overrides = {"physical_vacancy": {"value": 0.06, "note": "lease-up",
                                      "cohort_overrides": {"cohort_a": 0.05}}}
    cohort_ids = [c["cohort_id"] for c in resolved_deal["unit_cohorts"]]
    overrides["physical_vacancy"]["cohort_overrides"] = {cohort_ids[0]: 0.05}
    deal = resolve_assumptions(resolved_deal, overrides=overrides)
    path = write_assumptions_json(deal, tmp_path / "a.json")
    read = read_assumption_overrides(path)
    assert read["physical_vacancy"] == overrides["physical_vacancy"]
    again = resolve_assumptions(deal, overrides=read)
    assert again["assumptions"]["entries"]["physical_vacancy"] == deal["assumptions"]["entries"]["physical_vacancy"]


def test_plain_overrides_json_passes_through(tmp_path):
    path = tmp_path / "overrides.json"
    path.write_text(json.dumps({"rent_growth": 0.025}), encoding="utf-8")
    assert read_assumption_overrides(path) == {"rent_growth": 0.025}
    assert load_assumption_overrides(path) == {"rent_growth": 0.025}


@pytest.mark.parametrize("rows, message", [
    ([{"value": 1}], "no key"),
    ([{"key": "a"}, {"key": "a"}], "twice"),
    (["not a row"], "must be an object"),
])
def test_malformed_sheets_are_rejected(tmp_path, rows, message):
    path = tmp_path / "bad.json"
    path.write_text(json.dumps({"rows": rows}), encoding="utf-8")
    with pytest.raises(ValueError, match=message):
        read_assumption_overrides(path)


def test_json_edit_re_resolves_and_keeps_the_default(resolved_deal, tmp_path):
    path = write_assumptions_json(resolved_deal, tmp_path / "deal.assumptions.json")
    document = json.loads(path.read_text(encoding="utf-8"))
    row = next(r for r in document["rows"] if r["key"] == "rent_growth")
    default = row["default_value"]
    row["value"] = 0.025
    path.write_text(json.dumps(document), encoding="utf-8")

    deal = resolve_assumptions(resolved_deal, overrides=read_assumption_overrides(path))
    entry = deal["assumptions"]["entries"]["rent_growth"]
    assert entry["value"] == 0.025
    assert entry["override"] == {"value": 0.025, "note": DEFAULT_EDIT_NOTE}
    assert entry["default_value"] == default
    assert entry["source"] == "analyst"
    assert deal["growth_assumptions"]["annual_growth_rate"] == 0.025
    re_row = next(r for r in sheet_rows(deal) if r["key"] == "rent_growth")
    assert re_row["delta"] == pytest.approx(0.025 - default)


# ── Excel ────────────────────────────────────────────────────────────────


def _cell_for(path: Path, key: str, header: str):
    wb = load_workbook(path)
    ws = wb[SHEET_TITLE]
    table = ws.tables[TABLE_NAME]
    cells = ws[table.ref]
    headers = [c.value for c in cells[0]]
    for row in cells[1:]:
        if row[0].value == key:
            return wb, ws, row[headers.index(header)]
    raise AssertionError(f"{key} not in sheet")


def test_excel_round_trip_preserves_every_row_shape(synthetic_canonical, tmp_path):
    path = write_assumptions_xlsx(synthetic_canonical, tmp_path / "sheet.xlsx")
    wb = load_workbook(path)
    assert wb.sheetnames == [SHEET_TITLE]
    ws = wb[SHEET_TITLE]
    _, _, rent_value = _cell_for(path, "rent_growth", "Value")
    assert rent_value.value == "[0.021, 0.028, 0.03]"
    _, _, loss_value = _cell_for(path, "collection_loss", "Value")
    assert loss_value.value == 0.01 and loss_value.number_format == "0.00%"
    _, _, payroll = _cell_for(path, "opex.payroll.base", "Default Value")
    assert payroll.value == 1650.0 and payroll.number_format == '"$"#,##0.00'
    assert ws.freeze_panes == "B5"

    assert read_assumption_rows_xlsx(path) == sheet_rows(synthetic_canonical)
    assert read_assumption_overrides_xlsx(path) == {
        "collection_loss": {"value": 0.01, "note": "T12 bad debt"},
        "physical_vacancy": {"value": [0.08, 0.06, 0.05], "note": "slow lease-up"},
        "opex.payroll.base": {"value": 1800.0, "note": "new payroll plan"},
    }


def test_excel_edit_override_cell_read_back_and_re_resolve(resolved_deal, tmp_path):
    path = write_assumptions_xlsx(resolved_deal, tmp_path / "deal.xlsx")
    wb, ws, cell = _cell_for(path, "sale_cost_percent", "Override")
    note = ws.cell(row=cell.row, column=cell.column + 1)
    assert cell.value is None and note.value is None
    cell.value = 0.025
    note.value = "broker quote"
    wb.save(path)

    overrides = read_assumption_overrides_xlsx(path)
    expected = _existing_overrides(resolved_deal)
    expected["sale_cost_percent"] = {"value": 0.025, "note": "broker quote"}
    assert overrides == expected

    before = resolved_deal["assumptions"]["entries"]["sale_cost_percent"]
    deal = resolve_assumptions(resolved_deal, overrides=overrides)
    entry = deal["assumptions"]["entries"]["sale_cost_percent"]
    assert entry["value"] == 0.025
    assert entry["override"] == {"value": 0.025, "note": "broker quote"}
    assert entry["default_value"] == before["default_value"]
    assert deal["exit_assumptions"]["sale_cost_percent"] == 0.025
    # Every other entry is unchanged by the round trip.
    for key, other in resolved_deal["assumptions"]["entries"].items():
        if key != "sale_cost_percent":
            assert deal["assumptions"]["entries"][key] == other, key


def test_excel_value_edit_and_text_cells_are_decoded(resolved_deal, tmp_path):
    path = write_assumptions_xlsx(resolved_deal, tmp_path / "deal.xlsx")
    wb, ws, cell = _cell_for(path, "rent_growth", "Value")
    cell.value = "2.5%"
    wb.save(path)
    assert read_assumption_overrides_xlsx(path)["rent_growth"] == {"value": 0.025, "note": DEFAULT_EDIT_NOTE}


def test_excel_reader_falls_back_to_the_tab_when_the_table_is_dropped(synthetic_canonical, tmp_path):
    path = write_assumptions_xlsx(synthetic_canonical, tmp_path / "sheet.xlsx")
    wb = load_workbook(path)
    del wb[SHEET_TITLE].tables[TABLE_NAME]
    wb.save(path)
    assert read_assumption_rows_xlsx(path) == sheet_rows(synthetic_canonical)


def test_excel_with_no_entries_still_has_a_valid_table(tmp_path):
    canonical = {"assumptions": {"entries": {}}}
    path = write_assumptions_xlsx(canonical, tmp_path / "empty.xlsx")
    assert read_assumption_overrides_xlsx(path) == {}


def test_excel_invalid_json_cell_names_the_key(synthetic_canonical, tmp_path):
    path = write_assumptions_xlsx(synthetic_canonical, tmp_path / "sheet.xlsx")
    wb, _, cell = _cell_for(path, "rent_growth", "Override")
    cell.value = "[0.02, "
    wb.save(path)
    with pytest.raises(ValueError, match="rent_growth"):
        read_assumption_overrides_xlsx(path)


def test_underwriting_workbook_carries_the_assumptions_tab(minimal_deal_inputs, tmp_path):
    from engine.engine import run_underwriting
    from engine.excel_output import write_results_to_workbook

    inputs = resolve_assumptions(minimal_deal_inputs)
    results = run_underwriting(copy.deepcopy(inputs))
    path = tmp_path / "model.xlsx"
    write_results_to_workbook(results, path, inputs)
    assert SHEET_TITLE in load_workbook(path).sheetnames
    assert _norm(read_assumption_rows_xlsx(path)) == _norm(sheet_rows(inputs))

    plain = tmp_path / "plain.xlsx"
    write_results_to_workbook(results, plain, minimal_deal_inputs)
    assert SHEET_TITLE not in load_workbook(plain).sheetnames


def test_add_assumptions_sheet_replaces_an_existing_tab(synthetic_canonical):
    from openpyxl import Workbook

    wb = Workbook()
    add_assumptions_sheet(wb, synthetic_canonical)
    add_assumptions_sheet(wb, synthetic_canonical)
    assert wb.sheetnames.count(SHEET_TITLE) == 1


# ── run artifacts ────────────────────────────────────────────────────────


def test_run_from_files_writes_the_sheet(minimal_deal_inputs, tmp_path):
    from engine.engine import run_from_files

    inputs_path = tmp_path / "inputs.json"
    inputs_path.write_text(json.dumps(resolve_assumptions(minimal_deal_inputs)), encoding="utf-8")
    run_from_files(inputs_path, tmp_path / "outputs.json", tmp_path / "validation.json")
    sheet = tmp_path / "outputs.assumptions.json"
    assert sheet == sheet_path_for(tmp_path / "outputs.json")
    assert json.loads(sheet.read_text(encoding="utf-8"))["sheet_version"] == SHEET_VERSION

    plain = tmp_path / "plain_inputs.json"
    plain.write_text(json.dumps(minimal_deal_inputs), encoding="utf-8")
    run_from_files(plain, tmp_path / "plain.json", tmp_path / "plain_validation.json")
    assert not (tmp_path / "plain.assumptions.json").exists()


def test_api_run_response_carries_the_sheet(minimal_deal_inputs):
    from engine.api import handle_run_deal

    inputs = resolve_assumptions(minimal_deal_inputs)
    response = handle_run_deal({"inputs": inputs, "options": {"no_persist": True}})
    assert response["status"] == "success", response.get("error")
    assert response["assumptions_sheet"] == assumptions_document(inputs)

    plain = handle_run_deal({"inputs": minimal_deal_inputs, "options": {"no_persist": True}})
    assert plain["status"] == "success"
    assert "assumptions_sheet" not in plain


def test_backsolve_cli_writes_the_sheet_with_its_artifacts(minimal_deal_inputs, tmp_path, monkeypatch):
    from engine import backsolve

    case = resolve_assumptions(minimal_deal_inputs)
    canonical = tmp_path / "canonical.json"
    canonical.write_text(json.dumps(case), encoding="utf-8")
    monkeypatch.setattr(backsolve, "backsolve_price", lambda *a, **k: {
        "status": "converged", "case": case, "results": {}, "summary": {"status": "converged"},
    })
    monkeypatch.setattr(backsolve.sys, "argv", [
        "backsolve", "--canonical-json", str(canonical), "--output-dir", str(tmp_path / "out"),
        "--policy-version", backsolve.POLICY_VERSION, "--benchmark-5yr-treasury", "0.04",
        "--benchmark-as-of", "2026-01-02", "--benchmark-source", "synthetic",
    ])
    backsolve.main()
    summary = json.loads((tmp_path / "out" / "backsolve_summary.json").read_text(encoding="utf-8"))
    sheet = tmp_path / "out" / "assumptions.json"
    assert summary["artifacts"]["assumptions"] == str(sheet)
    assert json.loads(sheet.read_text(encoding="utf-8")) == assumptions_document(case)


# ── CLIs ─────────────────────────────────────────────────────────────────


def _ingest_argv(output: Path, *extra: str) -> list[str]:
    return [
        "ingest_deal.py", "--property-id", "Synthetic CLI", "--rent-roll", str(ROLL), "--t12", str(T12),
        "--start", "2026-08", "--end", "2031-07", "--output", str(output), *extra,
    ]


def test_ingest_cli_writes_the_sheet_next_to_the_output(tmp_path, monkeypatch, capsys):
    from runs import ingest_deal

    output = tmp_path / "deal.json"
    monkeypatch.setattr(sys, "argv", _ingest_argv(output))
    assert ingest_deal.main() == 0
    sheet = tmp_path / "deal.assumptions.json"
    assert f"Assumptions sheet: {sheet}" in capsys.readouterr().out
    deal = json.loads(output.read_text(encoding="utf-8"))
    assert json.loads(sheet.read_text(encoding="utf-8")) == assumptions_document(deal)


def test_ingest_cli_accepts_an_edited_sheet_as_overrides(tmp_path, monkeypatch, capsys):
    from runs import ingest_deal

    first = tmp_path / "first.json"
    monkeypatch.setattr(sys, "argv", _ingest_argv(first))
    assert ingest_deal.main() == 0
    deal = json.loads(first.read_text(encoding="utf-8"))

    # JSON sheet: edit a value.
    sheet = tmp_path / "first.assumptions.json"
    document = json.loads(sheet.read_text(encoding="utf-8"))
    next(r for r in document["rows"] if r["key"] == "collection_loss")["value"] = 0.0125
    sheet.write_text(json.dumps(document), encoding="utf-8")
    second = tmp_path / "second.json"
    monkeypatch.setattr(sys, "argv", _ingest_argv(second, "--assumption-overrides", str(sheet)))
    assert ingest_deal.main() == 0
    entry = json.loads(second.read_text(encoding="utf-8"))["assumptions"]["entries"]["collection_loss"]
    assert entry["value"] == 0.0125
    assert entry["default_value"] == deal["assumptions"]["entries"]["collection_loss"]["default_value"]

    # Excel sheet in --canonical-json mode: set an override cell.
    workbook = write_assumptions_xlsx(deal, tmp_path / "first.xlsx")
    wb, _, cell = _cell_for(workbook, "rent_growth", "Override")
    cell.value = 0.022
    wb.save(workbook)
    third = tmp_path / "third.json"
    monkeypatch.setattr(sys, "argv", [
        "ingest_deal.py", "--canonical-json", str(first), "--assumption-overrides", str(workbook),
        "--output", str(third),
    ])
    assert ingest_deal.main() == 0
    out = json.loads(third.read_text(encoding="utf-8"))
    rent = out["assumptions"]["entries"]["rent_growth"]
    assert rent["override"] == {"value": 0.022, "note": DEFAULT_EDIT_NOTE}
    assert rent["default_value"] == deal["assumptions"]["entries"]["rent_growth"]["default_value"]
    assert out["metadata"]["ingest_gate"]["assumed_defaults"] == [
        d for d in deal["metadata"]["ingest_gate"]["assumed_defaults"]
        if d["field"] != "growth_assumptions.annual_growth_rate"
    ]
    assert (tmp_path / "third.assumptions.json").exists()


def test_ingest_cli_rejects_an_unreadable_overrides_workbook(tmp_path, monkeypatch, capsys):
    from openpyxl import Workbook
    from runs import ingest_deal

    workbook = tmp_path / "empty.xlsx"
    Workbook().save(workbook)
    monkeypatch.setattr(sys, "argv", _ingest_argv(tmp_path / "deal.json", "--assumption-overrides", str(workbook)))
    assert ingest_deal.main() == 1
    assert "assumption overrides" in capsys.readouterr().err


def test_sheet_cli_export_and_apply(resolved_deal, tmp_path, capsys):
    from runs import assumptions_sheet

    canonical = tmp_path / "deal.json"
    canonical.write_text(json.dumps(resolved_deal), encoding="utf-8")

    assert assumptions_sheet.main(["export", str(canonical)]) == 0
    default_json = tmp_path / "deal.assumptions.json"
    assert json.loads(default_json.read_text(encoding="utf-8")) == assumptions_document(resolved_deal)

    workbook = tmp_path / "edit" / "deal.xlsx"
    json_out = tmp_path / "edit" / "sheet.json"
    assert assumptions_sheet.main(["export", str(canonical), "--xlsx", str(workbook), "--json", str(json_out)]) == 0
    out = capsys.readouterr().out
    assert str(workbook) in out and str(json_out) in out

    wb, ws, cell = _cell_for(workbook, "exit_cap_rate", "Override")
    cell.value = 0.0575
    ws.cell(row=cell.row, column=cell.column + 1).value = "IC exit view"
    wb.save(workbook)

    output = tmp_path / "applied.json"
    assert assumptions_sheet.main(["apply", str(canonical), str(workbook), "--output", str(output)]) == 0
    assert "Assumptions sheet:" in capsys.readouterr().out
    applied = json.loads(output.read_text(encoding="utf-8"))
    entry = applied["assumptions"]["entries"]["exit_cap_rate"]
    assert entry["override"] == {"value": 0.0575, "note": "IC exit view"}
    assert entry["default_value"] == resolved_deal["assumptions"]["entries"]["exit_cap_rate"]["default_value"]
    assert applied["exit_assumptions"]["exit_cap_rate"] == 0.0575
    assert (tmp_path / "applied.assumptions.json").exists()

    # JSON sheets apply the same way.
    document = json.loads(json_out.read_text(encoding="utf-8"))
    next(r for r in document["rows"] if r["key"] == "exit_cap_rate")["override_value"] = 0.0575
    json_out.write_text(json.dumps(document), encoding="utf-8")
    output_json = tmp_path / "applied_json.json"
    assert assumptions_sheet.main(["apply", str(canonical), str(json_out), "--output", str(output_json)]) == 0
    entry_json = json.loads(output_json.read_text(encoding="utf-8"))["assumptions"]["entries"]["exit_cap_rate"]
    assert entry_json["value"] == 0.0575


def test_sheet_cli_errors(minimal_deal_inputs, tmp_path, capsys):
    from runs import assumptions_sheet

    plain = tmp_path / "plain.json"
    plain.write_text(json.dumps(minimal_deal_inputs), encoding="utf-8")
    assert assumptions_sheet.main(["export", str(plain)]) == 2
    assert "no assumptions section" in capsys.readouterr().err

    assert assumptions_sheet.main(["apply", str(plain), str(tmp_path / "missing.json"),
                                   "--output", str(tmp_path / "o.json")]) == 1

    bad = tmp_path / "bad.json"
    bad.write_text(json.dumps({"not_a_key": 1}), encoding="utf-8")
    assert assumptions_sheet.main(["apply", str(plain), str(bad), "--output", str(tmp_path / "o.json")]) == 2
    assert "unknown assumption keys" in capsys.readouterr().err


def test_sheet_cli_runs_as_a_script(resolved_deal, tmp_path):
    canonical = tmp_path / "deal.json"
    canonical.write_text(json.dumps(resolved_deal), encoding="utf-8")
    result = subprocess.run(
        [str(VENV_PYTHON), "runs/assumptions_sheet.py", "export", str(canonical),
         "--xlsx", str(tmp_path / "deal.xlsx")],
        capture_output=True, text=True, cwd=str(REPO_ROOT),
    )
    assert result.returncode == 0, result.stderr
    assert read_assumption_overrides_xlsx(tmp_path / "deal.xlsx") == _existing_overrides(resolved_deal)


# ── determinism ──────────────────────────────────────────────────────────


def test_row_order_and_json_bytes_are_deterministic(synthetic_canonical, resolved_deal, tmp_path):
    shuffled = copy.deepcopy(synthetic_canonical)
    entries = shuffled["assumptions"]["entries"]
    shuffled["assumptions"]["entries"] = {key: entries[key] for key in reversed(list(entries))}
    assert sheet_rows(shuffled) == sheet_rows(synthetic_canonical)

    from engine.assumptions import ASSUMPTION_KEYS

    keys = [row["key"] for row in sheet_rows(resolved_deal)]
    assert keys == [key for key in ASSUMPTION_KEYS if key in keys]

    first = write_assumptions_json(resolved_deal, tmp_path / "a.json").read_bytes()
    second = write_assumptions_json(copy.deepcopy(resolved_deal), tmp_path / "b.json").read_bytes()
    assert first == second
    assert read_assumption_rows_xlsx(write_assumptions_xlsx(shuffled, tmp_path / "s.xlsx")) == sheet_rows(
        synthetic_canonical
    )
