"""The assumption register (ADR 0003 §8): one view of every assumption, rendered
as JSON, a self-contained HTML page, the Excel "Assumptions" tab, the API and
the MCP tool.

The register shows each assumption's value, the resolver default and its
source, why that source is the default (entry basis plus the policy rationale
in ``engine/assumptions/rationale.py``), the candidates considered, the user's
edit (with optional ``edited_by`` / ``edited_at``) and its delta from the
default, and the guardrail and review flags. All data is synthetic.
"""

from __future__ import annotations

import copy
import json
import re
from functools import lru_cache
from html.parser import HTMLParser
from pathlib import Path

import pytest
from openpyxl import load_workbook

from engine.api import handle_assumption_register
from engine.assumptions import AssumptionOverrideError, resolve_assumptions
from engine.assumptions.rationale import (
    NO_RATIONALE,
    PREFIX_RATIONALES,
    SOURCE_RATIONALES,
    rationale_for,
    source_family,
)
from engine.assumptions.register import (
    FILTERS,
    GROUPS,
    REGISTER_VERSION,
    build_register,
    format_value,
    group_for,
    label_for,
)
from engine.assumptions.register_html import (
    CHIPS,
    register_html_path_for,
    render_register_html,
    write_register_html,
)
from engine.assumptions.registry import SOURCE_PATTERN, override_metadata
from engine.assumptions.sheet import (
    EXCEL_COLUMN_ORDER,
    SHEET_TITLE,
    TABLE_NAME,
    document_rows,
    read_assumption_overrides_xlsx,
    read_assumption_rows_xlsx,
    stamp_overrides,
    write_assumptions_xlsx,
)
from engine.underwriting_policy import VINTAGE_FRAMEWORK_2026
from tests import deal_library as dl

REPO_ROOT = Path(__file__).resolve().parent.parent
SLUG_MARKET = "c_value_add_1970s_reno"
SLUG_HOUSE = "b_stabilized_2010s_no_market"


@lru_cache(maxsize=None)
def _ingested(slug: str) -> str:
    return json.dumps(dl.ingest(dl.load_deal(slug)))


def _deal(slug: str = SLUG_MARKET) -> dict:
    return json.loads(_ingested(slug))


def _row(register: dict, key: str) -> dict:
    return next(row for row in register["rows"] if row["key"] == key)


def _entry(value, default, **extra):
    entry = {
        "value": value,
        "unit": "decimal_yoy",
        "source": "house:long_run_rent_growth",
        "basis": "synthetic basis",
        "as_of": None,
        "confidence": "medium",
        "candidates": [{"source": "house:long_run_rent_growth", "value": default, "basis": "house rate"}],
        "default_value": default,
        "override": None,
        "review_required": True,
    }
    entry.update(extra)
    return entry


@pytest.fixture()
def shapes_canonical():
    """Entries of every shape, plus keys and sources the register has never seen."""
    return {
        "metadata": {"deal_id": "Synthetic Register"},
        "deal_profile": {"deal_type": "value_add", "msa": {"name": "Synthetic MSA"}, "year_built": 1984},
        "assumptions": {
            "resolver_version": "plat.assumptions/1",
            "mode": "underwrite",
            "entries": {
                "exit_cap_rate": _entry(
                    0.0625, 0.06, unit="decimal", source="analyst", review_required=False,
                    override={"value": 0.0625, "note": "IC exit view", "edited_by": "Analyst A",
                              "edited_at": "2026-10-01"},
                    candidates=[
                        {"source": "analyst", "value": 0.0625, "basis": "IC exit view",
                         "default": {"source": "house:exit_cap_spread", "basis": "going-in + spread",
                                     "as_of": None, "confidence": "medium"}},
                        {"source": "house:exit_cap_spread", "value": 0.06, "basis": "going-in + spread"},
                    ],
                ),
                "rent_growth": _entry(
                    [0.021, 0.028, 0.03, 0.03, 0.03], [0.021, 0.028, 0.03, 0.03, 0.03],
                    unit="decimal_yoy_by_analysis_year", source="market:synthetic", review_required=False,
                    as_of="2026-09-30", confidence="high",
                ),
                "physical_vacancy": _entry(
                    {"start": 0.08, "stabilized": 0.065, "window_months": 24},
                    {"start": 0.08, "stabilized": 0.065, "window_months": 24},
                    unit="decimal_of_units", source="deal:rent_roll", review_required=False,
                    guardrail={"status": "below_floor", "action": "warned", "floor": 0.08,
                               "floor_source": "house:vintage_1980s", "original": 0.065},
                ),
                "concessions": _entry(
                    {"year1": 0.03, "stabilized": 0.01, "year1_months": 12},
                    {"year1": 0.02, "stabilized": 0.01, "year1_months": 12},
                    unit="decimal_of_rent", source="analyst", review_required=False,
                    override={"value": {"year1": 0.03, "stabilized": 0.01, "year1_months": 12},
                              "note": "lease-up specials"},
                    candidates=[{"source": "analyst", "value": {"year1": 0.03}, "basis": "lease-up specials",
                                 "default": {"source": "deal:t12", "basis": "T12 concessions / GPR"}}],
                ),
                "opex.payroll.base": _entry(
                    180000.0, 180000.0, unit="usd_per_year", source="deal:t12", review_required=False,
                    detail={"category_name": "Payroll", "per_unit_annual": 1800.0},
                ),
                # Keys and a source no code knows yet (e.g. from upcoming capex work).
                "capex.roof.base": _entry(
                    250000.0, 250000.0, unit="usd", source="house:economic_capex_schedule",
                    candidates=[{"source": "house:economic_capex_schedule", "value": 250000.0}],
                    schedule={"2027": 100000.0, "2028": 150000.0},
                ),
                "zz_custom_mix": _entry({"1br": 0.5, "2br": 0.5}, {"1br": 0.5, "2br": 0.5},
                                        unit="share_by_unit_type", source="engine_fallback",
                                        confidence="low"),
            },
        },
    }


# ── register content ─────────────────────────────────────────────────────


def test_register_rows_cover_every_entry_shape(shapes_canonical):
    register = build_register(shapes_canonical)
    assert register["register_version"] == REGISTER_VERSION
    assert register["deal_id"] == "Synthetic Register"
    assert register["deal_profile"]["deal_type"] == "value_add"
    assert {row["key"] for row in register["rows"]} == set(shapes_canonical["assumptions"]["entries"])
    json.dumps(register)  # JSON-serializable

    rent = _row(register, "rent_growth")
    assert rent["group"] == "Revenue" and rent["label"] == "Market rent growth"
    assert rent["value_display"] == "2.1% → 2.8% … 3.0% (5 yrs)"
    assert rent["source_family"] == "market" and rent["as_of"] == "2026-09-30"

    vacancy = _row(register, "physical_vacancy")
    assert vacancy["group"] == "Vacancy & Credit"
    assert vacancy["value_display"] == "start 8.0% → 6.5% over 24 mo"
    assert vacancy["guardrail"]["state"] == "warned"
    assert vacancy["guardrail"]["floor"] == 0.08 and vacancy["guardrail"]["original"] == 0.065
    assert "guardrail_warned" in vacancy["flags"] and "guardrail" in vacancy["filters"]

    payroll = _row(register, "opex.payroll.base")
    assert payroll["group"] == "Opex" and payroll["label"] == "Payroll: base (year 1)"
    assert payroll["value_display"] == "$180,000"
    assert payroll["extra"] == {"detail": {"category_name": "Payroll", "per_unit_annual": 1800.0}}

    roof = _row(register, "capex.roof.base")
    assert roof["group"] == "Capital/Capex"
    assert roof["label"] == "Roof: base (year 1)"
    assert roof["extra"] == {"schedule": {"2027": 100000.0, "2028": 150000.0}}
    assert roof["why"]["rationale"] == NO_RATIONALE and not roof["why"]["rationale_known"]
    assert "no_rationale" in roof["flags"]

    other = _row(register, "zz_custom_mix")
    assert other["group"] == "Profile/Other"
    assert other["value_display"] == "1br: 0.5; 2br: 0.5"
    assert other["source_family"] == "engine_fallback" and "low_confidence" in other["flags"]
    assert "house" in other["filters"]

    counts = register["summary"]
    assert counts["total"] == 7
    assert counts["from_deal_evidence"] == 2 and counts["from_market"] == 1
    assert counts["house_defaults"] == 1 and counts["engine_fallbacks"] == 1
    assert counts["edited_by_user"] == 2 and counts["guardrail_warnings"] == 1
    assert counts["missing_rationale"] == 1


def test_rows_are_grouped_in_group_order_then_key_order(shapes_canonical):
    register = build_register(shapes_canonical)
    groups = [row["group"] for row in register["rows"]]
    assert groups == sorted(groups, key=GROUPS.index)
    assert [g["group"] for g in register["groups"]] == [g for g in GROUPS if g in groups]
    assert [key for g in register["groups"] for key in g["keys"]] == [row["key"] for row in register["rows"]]


def test_group_and_label_fall_back_for_unknown_keys():
    assert group_for("economic_capex_per_unit") == "Capital/Capex"
    assert group_for("r_and_m_floor_per_unit") == "Opex"
    assert group_for("loan_rate_cap") == "Debt/Fees"
    assert group_for("brand_new_thing") == "Profile/Other"
    assert label_for("brand_new_thing") == "Brand new thing"
    assert label_for("opex.water_sewer.growth") == "Water sewer: growth"


@pytest.mark.parametrize(
    "value, unit, text",
    [
        (None, "decimal", "—"),
        (0.025, "decimal", "2.5%"),
        (0.0575, "decimal", "5.75%"),
        (1350.0, "usd_per_unit_year", "$1,350"),
        (12.5, "usd_per_unit_month", "$12.50"),
        ([0.025] * 5, "decimal_yoy_by_analysis_year", "2.5% → 2.5% … 2.5% (5 yrs)"),
        ([0.02, 0.03], "decimal_yoy_by_analysis_year", "2.0% → 3.0% (2 yrs)"),
        ({"start": 0.045, "stabilized": 0.0575, "window_months": 12}, "decimal", "start 4.5% → 5.75% over 12 mo"),
        ({"year1": 0.01, "stabilized": 0.012, "year1_months": 12}, "decimal", "year 1 1.0% → 1.2% after 12 mo"),
        ({"2026Q4": 0.02, "2027Q1": 0.022, "2027Q2": 0.025, "2027Q3": 0.028}, "decimal",
         "2026Q4 2.0% → 2027Q1 2.2% … 2027Q3 2.8% (4 values)"),
    ],
)
def test_values_render_compactly(value, unit, text):
    assert format_value(value, unit) == text


def test_overridden_entry_shows_its_default_source_and_the_edit(shapes_canonical):
    row = _row(build_register(shapes_canonical), "exit_cap_rate")
    assert row["source"] == "analyst"
    assert row["default"]["source"] == "house:exit_cap_spread"
    assert row["default"]["value"] == 0.06 and row["default"]["display"] == "6.0%"
    assert row["why"]["rationale"] == SOURCE_RATIONALES["house:exit_cap_spread"]
    assert row["why"]["default_basis"] == "going-in + spread"
    edit = row["user_edit"]
    assert edit["value"] == 0.0625 and edit["note"] == "IC exit view"
    assert edit["edited_by"] == "Analyst A" and edit["edited_at"] == "2026-10-01"
    assert edit["differs_from_default"] is True
    assert edit["delta"] == pytest.approx(0.0025) and edit["delta_display"] == "+25 bps"
    assert row["flags"][0] == "edited" and "edited" in row["filters"]
    assert [c["source"] for c in row["candidates"]] == ["analyst", "house:exit_cap_spread"]
    assert row["candidates"][1]["is_default_source"] is True


def test_curve_edit_reports_a_per_field_delta(shapes_canonical):
    edit = _row(build_register(shapes_canonical), "concessions")["user_edit"]
    assert edit["differs_from_default"] is True
    assert edit["delta"] == {"stabilized": 0.0, "year1": pytest.approx(0.01), "year1_months": 0.0}
    assert edit["delta_display"] == "year1 +100 bps"


def test_filter_and_bounded_rows(shapes_canonical):
    full = build_register(shapes_canonical)
    edited = build_register(shapes_canonical, filter="edited", include_candidates=False)
    assert [row["key"] for row in edited["rows"]] == ["concessions", "exit_cap_rate"]
    assert edited["summary"] == full["summary"]  # the summary always counts every entry
    assert all("candidates" not in row and row["candidate_count"] >= 1 for row in edited["rows"])
    with pytest.raises(ValueError):
        build_register(shapes_canonical, filter="nope")
    assert set(FILTERS) == {chip for chip, _, _ in CHIPS}


def test_register_of_a_library_deal(shapes_canonical):
    deal = _deal(SLUG_MARKET)
    register = build_register(deal)
    entries = deal["assumptions"]["entries"]
    summary = register["summary"]
    assert summary["total"] == len(entries)
    assert summary["needing_review"] == sum(1 for e in entries.values() if e["review_required"])
    assert summary["guardrail_warnings"] == sum(
        1 for e in entries.values() if (e.get("guardrail") or {}).get("action") == "warned"
    )
    assert summary["guardrail_warnings"] >= 1
    assert summary["from_deal_evidence"] + summary["from_market"] + summary["house_defaults"] + summary[
        "engine_fallbacks"] + summary["analyst"] == summary["total"]
    assert summary["missing_rationale"] == 0
    assert register["deal_profile"] == deal["deal_profile"]
    for row in register["rows"]:
        assert row["value"] == entries[row["key"]]["value"]
        assert row["why"]["rationale_known"], row["key"]


# ── rationale coverage ───────────────────────────────────────────────────

_SOURCE_LITERAL = re.compile(
    r"""["'](analyst|engine_fallback|deal:[a-z0-9_]+|market:[A-Za-z0-9_.\-]+|house:[a-z0-9_]+)["']"""
)
_MARKET_SOURCE_CALL = re.compile(r"""market_source\(\s*["']([A-Za-z0-9_.\-]+)["']\s*\)""")
_SCANNED = sorted((REPO_ROOT / "engine" / "assumptions").glob("*.py")) + [REPO_ROOT / "engine" / "backsolve.py"]


def _sources_in_code() -> set[str]:
    found: set[str] = set()
    for path in _SCANNED:
        if path.name == "rationale.py":
            continue
        text = path.read_text(encoding="utf-8")
        found.update(_SOURCE_LITERAL.findall(text))
        found.update(f"market:{provider}" for provider in _MARKET_SOURCE_CALL.findall(text))
        if "house_vintage_source(" in text:
            found.update(f"house:vintage_{bucket}" for bucket in VINTAGE_FRAMEWORK_2026)
    return found


def test_every_source_the_code_emits_has_a_rationale():
    sources = _sources_in_code()
    # The grep finds the vocabulary the ADR lists.
    assert {"analyst", "engine_fallback", "deal:t12", "deal:rent_roll", "deal:om", "deal:comps",
            "house:backsolve_policy", "house:exit_cap_spread", "house:opex_growth_table",
            "house:long_run_rent_growth", "house:vintage_1980s", "market:supplied_benchmark"} <= sources
    missing = sorted(source for source in sources if not rationale_for(source)["known"])
    assert missing == [], (
        f"sources without a policy rationale: {missing}; add them to engine/assumptions/rationale.py"
    )


def test_rationale_vocabulary_is_valid_and_written_for_the_reader():
    for source, text in SOURCE_RATIONALES.items():
        assert SOURCE_PATTERN.match(source), source
        sentences = [s for s in re.split(r"(?<=[.!?])\s+", text.strip()) if s]
        assert 1 <= len(sentences) <= 3, source
    for prefix, template in PREFIX_RATIONALES:
        assert "{suffix}" in template, prefix
    vintage = rationale_for("house:vintage_1970s")
    assert vintage["known"] and "1970s" in vintage["text"] and vintage["matched"] == "house:vintage_"
    provider = rationale_for("market:some_new_feed")
    assert provider["known"] and provider["family"] == "market" and "some_new_feed" in provider["text"]
    for unknown in ("house:not_written_yet", "deal:new_document", None, "", "weird"):
        result = rationale_for(unknown)
        assert result["text"] == NO_RATIONALE and result["known"] is False
    assert source_family("house:vintage_1970s") == "house"
    assert source_family("engine_fallback") == "engine_fallback"
    assert source_family("nonsense") == "unknown"


# ── edits and edit metadata ──────────────────────────────────────────────


def test_edited_rows_are_flagged_with_their_delta():
    deal = _deal(SLUG_MARKET)
    before = deal["assumptions"]["entries"]["exit_cap_rate"]["value"]
    edited = resolve_assumptions(deal, overrides={
        "exit_cap_rate": {"value": round(before + 0.005, 6), "note": "IC view"},
        "rent_growth": [0.02] * len(deal["assumptions"]["entries"]["rent_growth"]["value"]),
    })
    register = build_register(edited)
    exit_row = _row(register, "exit_cap_rate")
    assert exit_row["user_edit"]["differs_from_default"] is True
    assert exit_row["user_edit"]["delta"] == pytest.approx(0.005)
    assert exit_row["user_edit"]["delta_display"] == "+50 bps"
    assert exit_row["default"]["source"].startswith("market:")
    rent_row = _row(register, "rent_growth")
    default = rent_row["default"]["value"]
    assert rent_row["user_edit"]["delta"] == pytest.approx([0.02 - d for d in default])
    # A recorded analyst input equal to its default is an edit that does not differ.
    sale = _row(register, "sale_cost_percent")
    assert sale["user_edit"]["differs_from_default"] is False and sale["user_edit"]["delta"] is None
    assert register["summary"]["edits_differing_from_default"] == 2


def test_override_metadata_round_trips_without_changing_values():
    deal = _deal(SLUG_HOUSE)
    plain = resolve_assumptions(deal, overrides={"collection_loss": {"value": 0.02, "note": "bad debt"}})
    tagged = resolve_assumptions(deal, overrides={"collection_loss": {
        "value": 0.02, "note": "bad debt", "edited_by": "  Analyst   B ", "edited_at": "2026-10-01T09:30:00",
    }})
    override = tagged["assumptions"]["entries"]["collection_loss"]["override"]
    assert override == {"value": 0.02, "note": "bad debt", "edited_by": "Analyst B", "edited_at": "2026-10-01"}

    # Values are the same with or without metadata.
    stripped = copy.deepcopy(tagged)
    for name in ("edited_by", "edited_at"):
        stripped["assumptions"]["entries"]["collection_loss"]["override"].pop(name)
    assert stripped == plain

    # Re-resolution (with no overrides) keeps the metadata and is idempotent.
    again = resolve_assumptions(tagged)
    assert again == tagged
    assert resolve_assumptions(again) == again

    row = _row(build_register(again), "collection_loss")
    assert row["user_edit"]["edited_by"] == "Analyst B" and row["user_edit"]["edited_at"] == "2026-10-01"

    # Clearing the override drops it, metadata included.
    cleared = resolve_assumptions(tagged, overrides={"collection_loss": None})
    assert cleared["assumptions"]["entries"]["collection_loss"]["override"] is None


@pytest.mark.parametrize("bad", [{"edited_by": ""}, {"edited_by": 3}, {"edited_at": "yesterday"},
                                 {"edited_at": "2026-13-01"}, {"edited_at": 20261001}, {"edited_by": "x" * 200}])
def test_bad_override_metadata_is_refused(bad):
    with pytest.raises(AssumptionOverrideError):
        resolve_assumptions(_deal(SLUG_HOUSE), overrides={"collection_loss": {"value": 0.02, **bad}})
    with pytest.raises(ValueError):
        override_metadata(bad)


def test_stamp_overrides_marks_only_new_or_changed_edits():
    deal = resolve_assumptions(_deal(SLUG_HOUSE), overrides={"collection_loss": {
        "value": 0.02, "note": "bad debt", "edited_by": "Analyst B", "edited_at": "2026-10-01"}})
    overrides = {
        "collection_loss": {"value": 0.02, "note": "bad debt"},       # as recorded
        "exit_cap_rate": {"value": 0.0611, "note": "IC exit view"},   # new edit
        "sale_cost_percent": None,                                    # clear
    }
    stamped = stamp_overrides(overrides, deal, edited_by="Analyst C", edited_at="2026-10-12")
    assert stamped["collection_loss"]["edited_by"] == "Analyst B"
    assert stamped["collection_loss"]["edited_at"] == "2026-10-01"
    assert stamped["exit_cap_rate"]["edited_by"] == "Analyst C"
    assert stamped["exit_cap_rate"]["edited_at"] == "2026-10-12"
    assert stamped["sale_cost_percent"] is None
    assert overrides["exit_cap_rate"] == {"value": 0.0611, "note": "IC exit view"}  # input untouched
    # Without metadata to add, overrides pass through unchanged.
    fresh = _deal(SLUG_HOUSE)
    assert stamp_overrides({"exit_cap_rate": 0.06, "collection_loss": {"value": 0.02}}, fresh) == {
        "exit_cap_rate": 0.06, "collection_loss": {"value": 0.02}}


# ── HTML ─────────────────────────────────────────────────────────────────

_VOID = {"meta", "br", "col", "img", "hr", "input", "link", "wbr", "source", "area", "base", "embed", "track"}


class _Checker(HTMLParser):
    """Balanced-tag checker: every non-void element is closed in order."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.stack: list[str] = []
        self.errors: list[str] = []
        self.tags: list[str] = []
        self.text: list[str] = []

    def handle_starttag(self, tag, attrs):
        self.tags.append(tag)
        if tag not in _VOID:
            self.stack.append(tag)

    def handle_startendtag(self, tag, attrs):
        self.tags.append(tag)
        if tag not in _VOID:
            self.errors.append(f"self-closed non-void <{tag}/>")

    def handle_endtag(self, tag):
        if tag in _VOID:
            return
        if not self.stack or self.stack[-1] != tag:
            self.errors.append(f"</{tag}> closes {self.stack[-1] if self.stack else 'nothing'}")
            return
        self.stack.pop()

    def handle_data(self, data):
        self.text.append(data)


def _check(html: str) -> _Checker:
    checker = _Checker()
    checker.feed(html)
    checker.close()
    assert checker.errors == [], checker.errors[:5]
    assert checker.stack == [], checker.stack
    return checker


def _hostile(canonical: dict) -> dict:
    hostile = copy.deepcopy(canonical)
    hostile["metadata"]["deal_id"] = 'Deal "><b>bold</b>'
    entry = hostile["assumptions"]["entries"]["exit_cap_rate"]
    entry["override"]["note"] = "<script>alert('x')</script> & more"
    entry["override"]["edited_by"] = '<img src=x onerror="alert(1)">'
    entry["basis"] = "</td></tr><tr><td>injected"
    hostile["assumptions"]["entries"]["<b>key</b>"] = _entry(0.01, 0.01, source="house:<i>x</i>")
    return hostile


def test_html_is_well_formed_and_self_contained(shapes_canonical):
    html = render_register_html(shapes_canonical)
    assert html.startswith("<!DOCTYPE html>")
    checker = _check(html)
    assert checker.tags.count("script") == 1 and checker.tags.count("style") == 1
    assert "http://" not in html and "https://" not in html  # no external assets
    assert not re.search(r"<(link|img|iframe)\b", html)
    assert "prefers-color-scheme: dark" in html and "@media print" in html
    for _, label, _ in CHIPS:
        assert f">{label}<" in html
    assert html.count('<tbody class="group"') == len({row["group"] for row in build_register(shapes_canonical)["rows"]})
    # Edited rows are highlighted; the edit shows its note, editor and date.
    assert re.search(r'<tr class="row edited[^"]*" id="a-exit_cap_rate"', html)
    assert "IC exit view" in html and "by Analyst A on 2026-10-01" in html
    assert "+25 bps vs default" in html
    assert "Candidates considered (2)" in html
    assert NO_RATIONALE in html


def test_html_escapes_every_injected_text(shapes_canonical):
    html = render_register_html(_hostile(shapes_canonical))
    checker = _check(html)
    assert checker.tags.count("script") == 1  # only the page's own filter script
    assert "<b>bold</b>" not in html and "<b>key</b>" not in html and "<i>x</i>" not in html
    assert "<script>alert" not in html and "<img" not in html and "</td></tr><tr><td>injected" not in html
    text = "".join(checker.text)
    assert "<script>alert('x')</script> & more" in text
    assert 'Deal "><b>bold</b>' in text
    assert "&lt;script&gt;alert(&#x27;x&#x27;)&lt;/script&gt; &amp; more" in html


def test_html_and_register_are_deterministic(shapes_canonical, tmp_path):
    shuffled = copy.deepcopy(shapes_canonical)
    entries = shuffled["assumptions"]["entries"]
    shuffled["assumptions"]["entries"] = {key: entries[key] for key in reversed(list(entries))}
    assert build_register(shuffled) == build_register(shapes_canonical)
    assert json.dumps(build_register(shuffled)) == json.dumps(build_register(shapes_canonical))
    assert render_register_html(shuffled) == render_register_html(shapes_canonical)

    deal = _deal(SLUG_MARKET)
    first = write_register_html(deal, tmp_path / "a.html").read_bytes()
    second = write_register_html(copy.deepcopy(deal), tmp_path / "b.html").read_bytes()
    assert first == second
    assert not re.search(rb"20\d\d-\d\d-\d\dT\d\d:\d\d", first)  # no generated timestamp
    _check(first.decode("utf-8"))


def test_register_html_path_sits_next_to_the_output():
    assert register_html_path_for(Path("out/deal.json")) == Path("out/deal.assumptions.html")


# ── Excel ────────────────────────────────────────────────────────────────


def _table_rows(path: Path):
    wb = load_workbook(path)
    ws = wb[SHEET_TITLE]
    cells = ws[ws.tables[TABLE_NAME].ref]
    headers = [c.value for c in cells[0]]
    return wb, ws, headers, cells[1:]


def test_excel_tab_shows_the_register_and_still_round_trips(tmp_path):
    deal = resolve_assumptions(_deal(SLUG_MARKET), overrides={"exit_cap_rate": {
        "value": 0.0625, "note": "IC exit view", "edited_by": "Analyst A", "edited_at": "2026-10-01"}})
    path = write_assumptions_xlsx(deal, tmp_path / "deal.xlsx")
    wb, ws, headers, body = _table_rows(path)
    for header in ("Group", "Default Source", "Why This Default", "Edited By / At", "Flags"):
        assert header in headers
    assert headers[0] == "Key" and headers.index("Note") == headers.index("Override") + 1
    assert len(headers) == len(EXCEL_COLUMN_ORDER)
    assert ws.freeze_panes == "B5"

    register = build_register(deal)
    keys = [row[0].value for row in body]
    assert keys == [row["key"] for row in register["rows"]]  # sorted by group, then key order
    groups = [row[headers.index("Group")].value for row in body]
    assert groups == sorted(groups, key=GROUPS.index)

    exit_row = next(row for row in body if row[0].value == "exit_cap_rate")
    assert exit_row[headers.index("Edited By / At")].value == "Analyst A · 2026-10-01"
    assert exit_row[headers.index("Default Source")].value.startswith("market:")
    assert exit_row[headers.index("Flags")].value.startswith("Edited by user")
    assert exit_row[headers.index("Group")].fill.start_color.rgb.endswith("FFF2CC")  # highlighted
    assert exit_row[0].font.bold
    review_row = next(row for row in body if row[headers.index("Review Required")].value is True)
    assert "Needs review" in review_row[headers.index("Flags")].value
    assert review_row[headers.index("Flags")].fill.start_color.rgb.endswith("FBE3E0")
    why = next(row for row in body if row[0].value == "rent_growth")[headers.index("Why This Default")].value
    assert why == rationale_for("market:synthetic")["text"]

    # The reader ignores the register columns and the display order.
    assert read_assumption_rows_xlsx(path) == document_rows(deal)
    cell = exit_row[headers.index("Override")]
    exit_row[headers.index("Group")].value = "edited by mistake"
    cell.value = 0.065
    wb.save(path)
    overrides = read_assumption_overrides_xlsx(path)
    assert overrides["exit_cap_rate"] == {"value": 0.065, "note": "IC exit view"}
    applied = resolve_assumptions(deal, overrides=stamp_overrides(overrides, deal, edited_by="Analyst C"))
    entry = applied["assumptions"]["entries"]["exit_cap_rate"]
    assert entry["value"] == 0.065 and entry["override"]["edited_by"] == "Analyst C"
    assert "edited_at" not in entry["override"]


# ── CLI ──────────────────────────────────────────────────────────────────


def test_sheet_cli_html_and_apply_with_editor(tmp_path, capsys, monkeypatch):
    from runs import assumptions_sheet

    canonical = tmp_path / "deal.json"
    canonical.write_text(_ingested(SLUG_MARKET), encoding="utf-8")

    assert assumptions_sheet.main(["html", str(canonical)]) == 0
    default_page = tmp_path / "deal.assumptions.html"
    assert f"Assumption register: {default_page}" in capsys.readouterr().out
    _check(default_page.read_text(encoding="utf-8"))
    custom = tmp_path / "pages" / "register.html"
    assert assumptions_sheet.main(["html", str(canonical), "-o", str(custom)]) == 0
    assert custom.read_bytes() == default_page.read_bytes()

    plain = tmp_path / "plain.json"
    plain.write_text(json.dumps({"metadata": {}}), encoding="utf-8")
    assert assumptions_sheet.main(["html", str(plain)]) == 2

    # Apply an edit with an editor and a date: only the changed row is stamped.
    overrides = tmp_path / "overrides.json"
    overrides.write_text(json.dumps({"exit_cap_rate": {"value": 0.0625, "note": "IC exit view"}}), encoding="utf-8")
    output = tmp_path / "applied.json"
    assert assumptions_sheet.main(["apply", str(canonical), str(overrides), "--output", str(output),
                                   "--no-market-data", "--editor", "Analyst A", "--edit-date", "2026-10-01"]) == 0
    out = capsys.readouterr().out
    assert f"Assumption register: {tmp_path / 'applied.assumptions.html'}" in out
    entries = json.loads(output.read_text(encoding="utf-8"))["assumptions"]["entries"]
    assert entries["exit_cap_rate"]["override"] == {
        "value": 0.0625, "note": "IC exit view", "edited_by": "Analyst A", "edited_at": "2026-10-01"}
    assert "edited_by" not in entries["sale_cost_percent"]["override"]  # an earlier input, not this edit
    page = (tmp_path / "applied.assumptions.html").read_text(encoding="utf-8")
    assert "by Analyst A on 2026-10-01" in page

    # Re-applying the exported sheet as another editor keeps the recorded editor.
    from engine.assumptions.sheet import write_assumptions_json

    sheet = write_assumptions_json(json.loads(output.read_text(encoding="utf-8")), tmp_path / "sheet.json")
    again = tmp_path / "again.json"

    class _Today:
        @staticmethod
        def today():
            import datetime as _dt
            return _dt.date(2026, 10, 12)

    monkeypatch.setattr(assumptions_sheet, "date", _Today)
    assert assumptions_sheet.main(["apply", str(output), str(sheet), "--output", str(again),
                                   "--no-market-data", "--editor", "Analyst B", "--stamp-date"]) == 0
    kept = json.loads(again.read_text(encoding="utf-8"))["assumptions"]["entries"]["exit_cap_rate"]["override"]
    assert kept["edited_by"] == "Analyst A" and kept["edited_at"] == "2026-10-01"

    document = json.loads(sheet.read_text(encoding="utf-8"))
    next(r for r in document["rows"] if r["key"] == "exit_cap_rate")["override_value"] = 0.064
    sheet.write_text(json.dumps(document), encoding="utf-8")
    assert assumptions_sheet.main(["apply", str(output), str(sheet), "--output", str(again),
                                   "--no-market-data", "--editor", "Analyst B", "--stamp-date"]) == 0
    changed = json.loads(again.read_text(encoding="utf-8"))["assumptions"]["entries"]["exit_cap_rate"]["override"]
    assert changed["value"] == 0.064
    assert changed["edited_by"] == "Analyst B" and changed["edited_at"] == "2026-10-12"


def test_run_deal_writes_the_register_page(tmp_path, capsys):
    from runs import run_deal

    canonical = tmp_path / "deal.json"
    canonical.write_text(_ingested(SLUG_HOUSE), encoding="utf-8")
    out = tmp_path / "run"
    assert run_deal.main([str(canonical), "--output-dir", str(out)]) == 0
    assert f"Assumption register: {out / 'assumptions.html'}" in capsys.readouterr().out
    page = (out / "assumptions.html").read_text(encoding="utf-8")
    _check(page)
    assert (out / "assumptions.json").exists()
    assert page == render_register_html(json.loads((out / "canonical.json").read_text(encoding="utf-8")))


def test_ingest_deal_writes_the_register_page(tmp_path, capsys, monkeypatch):
    import sys

    from runs import ingest_deal

    canonical = tmp_path / "deal.json"
    canonical.write_text(_ingested(SLUG_HOUSE), encoding="utf-8")
    output = tmp_path / "ingested.json"
    monkeypatch.setattr(sys, "argv", [
        "ingest_deal.py", "--canonical-json", str(canonical), "--output", str(output), "--no-market-data",
    ])
    assert ingest_deal.main() == 0
    page = tmp_path / "ingested.assumptions.html"
    assert f"Assumption register: {page}" in capsys.readouterr().out
    _check(page.read_text(encoding="utf-8"))
    assert (tmp_path / "ingested.assumptions.json").exists()


# ── API ──────────────────────────────────────────────────────────────────


def test_api_handler_returns_the_register(shapes_canonical):
    response = handle_assumption_register({"inputs": shapes_canonical})
    assert response["status"] == "success"
    assert response["deal_id"] == "Synthetic Register"
    assert response["register"] == build_register(shapes_canonical)
    assert "html" not in response
    json.dumps(response)

    filtered = handle_assumption_register({"inputs": shapes_canonical, "options": {
        "filter": "review", "include_candidates": False, "include_html": True}})
    assert filtered["status"] == "success"
    assert all(row["review_required"] for row in filtered["register"]["rows"])
    assert all("candidates" not in row for row in filtered["register"]["rows"])
    _check(filtered["html"])

    assert handle_assumption_register({})["status"] == "error"
    assert "assumptions section" in handle_assumption_register({"inputs": {"metadata": {}}})["error"]
    assert handle_assumption_register({"inputs": shapes_canonical, "options": {"filter": "x"}})["status"] == "error"
