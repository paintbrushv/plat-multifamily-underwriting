import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

from engine.excel_io import load_inputs_from_workbook
from engine.validator import validate_deal


def _add_table(ws, name: str, top_left: str, headers: list[str], rows: list[list]):
    start_col = ws[top_left].column
    start_row = ws[top_left].row
    for j, h in enumerate(headers):
        ws.cell(row=start_row, column=start_col + j, value=h)
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row):
            ws.cell(row=start_row + i, column=start_col + j, value=v)

    end_row = start_row + len(rows)
    end_col = start_col + len(headers) - 1
    ref = f"{ws.cell(row=start_row, column=start_col).coordinate}:{ws.cell(row=end_row, column=end_col).coordinate}"
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
    ws.add_table(t)


class TestExcelIO(unittest.TestCase):
    def test_load_inputs_from_workbook_happy_path(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "INPUTS_CORE"

        _add_table(
            ws,
            "tbl_metadata",
            "A1",
            ["deal_id", "run_id", "as_of_date", "analyst", "purpose", "notes"],
            [["d1", "r1", date(2026, 1, 1), "a", "Screening", ""]],
        )
        _add_table(ws, "tbl_time_grid", "A4", ["analysis_start_date", "analysis_end_date"], [[date(2026, 1, 1), date(2026, 1, 1)]])
        _add_table(
            ws,
            "tbl_unit_cohorts",
            "A7",
            ["cohort_id", "unit_type", "unit_count", "sqft", "initial_inplace_rent"],
            [["A", "1B", 10, 650, 1000]],
        )

        ws2 = wb.create_sheet("INPUTS_RENT")
        _add_table(ws2, "tbl_market_rent_curve", "A1", ["cohort_id", "start_period", "end_period", "market_rent"], [["A", "2026-01", "2026-01", 1200]])
        _add_table(ws2, "tbl_loss_to_lease", "A4", ["cohort_id", "start_period", "end_period", "ltl_percent"], [["A", "2026-01", "2026-01", 0.1]])
        _add_table(
            ws2,
            "tbl_physical_vacancy_curve",
            "A7",
            ["cohort_id", "start_period", "end_period", "vacancy_rate"],
            [["A", "2026-01", "2026-01", 0.2]],
        )
        _add_table(ws2, "tbl_collection_loss_curve", "A10", ["applies_to", "start_period", "end_period", "loss_rate"], [["ALL", "2026-01", "2026-01", 0.05]])

        ws3 = wb.create_sheet("INPUTS_PROGRAMS")
        _add_table(
            ws3,
            "tbl_revenue_programs",
            "A1",
            ["program_id", "program_name", "program_type", "pricing_type", "price_value", "eligible_units", "start_period", "end_period"],
            [["parking", "Parking", "tenant-based", "$/unit", 50, "ALL", "2026-01", "2026-01"]],
        )
        _add_table(ws3, "tbl_program_adoption_curve", "A4", ["program_id", "start_period", "end_period", "adoption_rate"], [["parking", "2026-01", "2026-01", 0.5]])
        _add_table(ws3, "tbl_program_capacity", "A7", ["program_id", "total_capacity"], [["parking", 6]])
        _add_table(ws3, "tbl_program_costs", "A10", ["program_id", "cost_type", "cost_value"], [["parking", "% revenue", 0.0]])

        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "deal.xlsx"
            wb.save(path)
            inputs = load_inputs_from_workbook(path)

        self.assertEqual(inputs["metadata"]["deal_id"], "d1")
        self.assertEqual(inputs["metadata"]["run_id"], "r1")
        self.assertEqual(inputs["time_grid"]["analysis_start_date"], "2026-01-01")
        self.assertEqual(inputs["unit_cohorts"][0]["unit_count"], 10)

        report = validate_deal(inputs)
        self.assertEqual(report.status, "PASS")


if __name__ == "__main__":
    unittest.main()

