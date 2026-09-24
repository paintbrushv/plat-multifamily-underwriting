"""
PowerPoint Presentation Generator
===================================
Creates a branded 5-slide LP presentation deck from engine results.

Slides:
1. Cover — Deal name, metro, units, company branding
2. Deal Summary — Key metrics table (purchase, cap rates, debt, returns)
3. Investment Returns — IRR/EM table + optional Bull/Base/Bear scenarios
4. Cash Flow — Annual NOI + Levered CF bar chart
5. Risk & Considerations — Key assumptions, sensitivities, risk factors

Default brand: Dark (#0a0a0a) + Accent Gold (#c9a66c)

Usage:
    from engine.pptx_generator import generate_presentation

    generate_presentation(results, inputs, "output/deal_deck.pptx")
    generate_presentation(results, inputs, "output/deal.pptx", scenario_results=scenarios)
"""
from __future__ import annotations

import io
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Optional

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.ticker as mticker

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor


# ── Brand Colors ──────────────────────────────────────────────────────────
from engine.brand import (
    MPL_DARK_900 as MPL_DARK, MPL_GOLD, MPL_GRAY, MPL_DARK_800 as MPL_BG,
    BrandConfig,
)
from engine.formatters import fmt_pct, fmt_currency, fmt_multiple, fmt_number

DARK_900 = RGBColor(0x0A, 0x0A, 0x0A)
DARK_800 = RGBColor(0x11, 0x11, 0x11)
DARK_700 = RGBColor(0x1A, 0x1A, 0x1A)
GOLD = RGBColor(0xC9, 0xA6, 0x6C)
GOLD_LIGHT = RGBColor(0xD4, 0xB8, 0x8A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xAA, 0xAA, 0xAA)
MEDIUM_GRAY = RGBColor(0x66, 0x66, 0x66)

# Slide dimensions (widescreen 16:9)
SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)


# ── Helpers ───────────────────────────────────────────────────────────────


def _set_slide_bg(slide, color: RGBColor = DARK_900):
    """Set slide background to solid color."""
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _add_textbox(slide, left, top, width, height, text, font_size=12,
                 font_name="Calibri", color=WHITE, bold=False,
                 alignment=PP_ALIGN.LEFT):
    """Add a text box to a slide."""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.name = font_name
    p.font.color.rgb = color
    p.font.bold = bold
    p.alignment = alignment
    return txBox


def _add_table(slide, left, top, width, height, rows, cols):
    """Add a table shape and return the table object."""
    table_shape = slide.shapes.add_table(rows, cols, left, top, width, height)
    return table_shape.table


def _style_header_row(table, col_count, bg=DARK_700, fg=GOLD):
    """Style the first row of a table as header."""
    for i in range(col_count):
        cell = table.cell(0, i)
        cell.fill.solid()
        cell.fill.fore_color.rgb = bg
        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.size = Pt(10)
            paragraph.font.bold = True
            paragraph.font.color.rgb = fg
            paragraph.font.name = "Calibri"
            paragraph.alignment = PP_ALIGN.CENTER


def _style_data_cell(cell, value, font_size=10, color=WHITE, bold=False,
                     alignment=PP_ALIGN.RIGHT, bg=None):
    """Style a data cell."""
    cell.text = str(value) if value is not None else "—"
    if bg:
        cell.fill.solid()
        cell.fill.fore_color.rgb = bg
    for paragraph in cell.text_frame.paragraphs:
        paragraph.font.size = Pt(font_size)
        paragraph.font.color.rgb = color
        paragraph.font.bold = bold
        paragraph.font.name = "Calibri"
        paragraph.alignment = alignment
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE


def _location_text(meta):
    address = meta.get("address")
    if isinstance(address, dict):
        city = str(address.get("city") or "").strip()
        state = str(address.get("state") or "").strip()
        if city and state:
            return f"{city}, {state}"
        street = str(address.get("street") or "").strip()
        if street:
            return street
    elif isinstance(address, str) and address.strip():
        parts = [part.strip() for part in address.split(",") if part.strip()]
        if len(parts) >= 2:
            return ", ".join(parts[-2:])
        return address.strip()
    market = meta.get("market")
    if market not in (None, ""):
        return str(market)
    city = meta.get("city")
    state = meta.get("state")
    if city and state:
        return f"{city}, {state}"
    return "—"


def _fmt_pct(val):
    return fmt_pct(val)

def _fmt_currency(val, decimals=0):
    return fmt_currency(val, decimals=decimals)

def _fmt_multiple(val):
    return fmt_multiple(val)

def _fmt_number(val):
    return fmt_number(val)


# ── Slide Builders ────────────────────────────────────────────────────────


def _build_cover_slide(prs, inputs, results, brand_config=None):
    """Slide 1: Cover with deal name, metro, key stats."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
    _set_slide_bg(slide)

    if brand_config is None:
        brand_config = BrandConfig()

    meta = inputs.get("metadata", {})
    deal_name = meta.get("deal_id", "Investment Opportunity")
    city = meta.get("city", "")
    state = meta.get("state", "")
    location = f"{city}, {state}" if city else ""

    cohorts = inputs.get("unit_cohorts", [])
    total_units = sum(c.get("unit_count", 0) for c in cohorts)
    pp = inputs.get("purchase_assumptions", {}).get("purchase_price", 0)

    # Gold accent bar at top
    shape = slide.shapes.add_shape(
        1,  # Rectangle
        Inches(0), Inches(0), SLIDE_WIDTH, Inches(0.08)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = GOLD
    shape.line.fill.background()

    # Partner logo on cover (top-right) if configured
    if brand_config.has_partner and brand_config.partner_logo_path:
        logo_path = brand_config.partner_logo_path
        if logo_path.exists():
            slide.shapes.add_picture(
                str(logo_path), Inches(10.5), Inches(0.5), Inches(2), Inches(0.8)
            )

    # Company name
    company_label = brand_config.company_name.upper()
    _add_textbox(slide, Inches(1), Inches(1.5), Inches(11), Inches(0.5),
                 company_label, font_size=14, color=GOLD,
                 font_name="Calibri", bold=True)

    # Deal name (large)
    _add_textbox(slide, Inches(1), Inches(2.2), Inches(11), Inches(1.2),
                 deal_name.upper(), font_size=44, color=WHITE,
                 font_name="Georgia", bold=True)

    # Location
    if location:
        _add_textbox(slide, Inches(1), Inches(3.5), Inches(11), Inches(0.5),
                     location, font_size=20, color=LIGHT_GRAY)

    # Key stats bar
    stats_y = Inches(4.5)
    stat_width = Inches(2.5)
    stats = [
        ("UNITS", _fmt_number(total_units)),
        ("PURCHASE PRICE", _fmt_currency(pp)),
    ]

    metrics = results.get("metrics", {})
    irr = metrics.get("irr", {}).get("levered_irr")
    em = metrics.get("equity_multiple", {}).get("levered_em")
    if irr is not None:
        stats.append(("LEVERED IRR", _fmt_pct(irr)))
    if em is not None:
        stats.append(("EQUITY MULTIPLE", _fmt_multiple(em)))

    for i, (label, value) in enumerate(stats):
        x = Inches(1) + Inches(3) * i

        # Label
        _add_textbox(slide, x, stats_y, stat_width, Inches(0.3),
                     label, font_size=10, color=GOLD, bold=True)

        # Value
        _add_textbox(slide, x, stats_y + Inches(0.35), stat_width, Inches(0.5),
                     value, font_size=24, color=WHITE, bold=True,
                     font_name="Georgia")

    # Bottom bar
    shape = slide.shapes.add_shape(
        1, Inches(0), Inches(7.42), SLIDE_WIDTH, Inches(0.08)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = GOLD
    shape.line.fill.background()

    # Confidential footer
    footer_label = _footer_text(brand_config)
    _add_textbox(slide, Inches(1), Inches(6.8), Inches(11), Inches(0.3),
                 footer_label,
                 font_size=8, color=MEDIUM_GRAY, alignment=PP_ALIGN.CENTER)

    # Date
    _add_textbox(slide, Inches(9), Inches(6.8), Inches(3.5), Inches(0.3),
                 datetime.now().strftime("%B %Y"),
                 font_size=8, color=MEDIUM_GRAY, alignment=PP_ALIGN.RIGHT)


def _footer_text(brand_config: BrandConfig) -> str:
    """Build confidential footer text, including partner name when co-branded."""
    if brand_config.has_partner:
        return f"CONFIDENTIAL — {brand_config.company_name.upper()} & {brand_config.partner_name.upper()}"
    return "CONFIDENTIAL — EXAMPLESPONSOR CAPITAL"


def _build_summary_slide(prs, inputs, results, brand_config=None):
    """Slide 2: Deal Summary — key metrics in a clean table layout."""
    if brand_config is None:
        brand_config = BrandConfig()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide)

    # Title
    _add_textbox(slide, Inches(0.8), Inches(0.4), Inches(11), Inches(0.6),
                 "DEAL SUMMARY", font_size=28, color=WHITE,
                 font_name="Georgia", bold=True)

    # Gold underline
    shape = slide.shapes.add_shape(
        1, Inches(0.8), Inches(1.05), Inches(2), Inches(0.04)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = GOLD
    shape.line.fill.background()

    meta = inputs.get("metadata", {})
    purchase = inputs.get("purchase_assumptions", {})
    exit_a = inputs.get("exit_assumptions", {})
    tg = inputs.get("time_grid", {})
    cohorts = inputs.get("unit_cohorts", [])
    metrics = results.get("metrics", {})
    irr_d = metrics.get("irr", {})
    em_d = metrics.get("equity_multiple", {})
    dscr_d = metrics.get("dscr", {})
    yields_d = metrics.get("yields", {})

    total_units = sum(c.get("unit_count", 0) for c in cohorts)
    total_sf = sum(c.get("unit_count", 0) * c.get("sqft", 0) for c in cohorts)
    pp = purchase.get("purchase_price", 0)

    # Two-column layout: left = property info, right = returns
    left_data = [
        ("Property", meta.get("deal_id", "—")),
        ("Location", _location_text(meta)),
        ("Total Units", _fmt_number(total_units)),
        ("Total SF", _fmt_number(total_sf)),
        ("Year Built", _fmt_number(meta.get("year_built"))),
        ("Purchase Price", _fmt_currency(pp)),
        ("Price / Unit", _fmt_currency(pp / total_units if total_units else 0)),
        ("Price / SF", _fmt_currency(pp / total_sf if total_sf else 0, 2)),
        ("Hold Period", f"{tg.get('analysis_start_date', '—')} to {tg.get('analysis_end_date', '—')}"),
        ("Going-In Cap", _fmt_pct(yields_d.get("going_in_cap_rate"))),
        ("Exit Cap", _fmt_pct(exit_a.get("exit_cap_rate"))),
    ]

    right_data = [
        ("Levered IRR", _fmt_pct(irr_d.get("levered_irr"))),
        ("Unlevered IRR", _fmt_pct(irr_d.get("unlevered_irr"))),
        ("Levered EM", _fmt_multiple(em_d.get("levered_em"))),
        ("Unlevered EM", _fmt_multiple(em_d.get("unlevered_em"))),
        ("Partnership IRR", _fmt_pct(irr_d.get("partnership_irr"))),
        ("Partnership EM", _fmt_multiple(em_d.get("partnership_em"))),
        ("Avg DSCR", f"{dscr_d.get('average_dscr', '—')}"),
        ("Min DSCR", f"{dscr_d.get('minimum_dscr', '—')}"),
        ("Cash-on-Cash Y1", _fmt_pct(
            metrics.get("cash_on_cash", {}).get("by_year", [{}])[0].get("yield")
            if metrics.get("cash_on_cash", {}).get("by_year") else None
        )),
    ]

    # Left table
    left_table = _add_table(
        slide, Inches(0.8), Inches(1.4), Inches(5.5), Inches(5),
        len(left_data), 2,
    )
    for i, (label, value) in enumerate(left_data):
        bg = DARK_800 if i % 2 == 0 else DARK_900
        _style_data_cell(left_table.cell(i, 0), label, color=GOLD,
                         bold=True, alignment=PP_ALIGN.LEFT, bg=bg)
        _style_data_cell(left_table.cell(i, 1), value, color=WHITE,
                         alignment=PP_ALIGN.RIGHT, bg=bg)

    left_table.columns[0].width = Inches(2.5)
    left_table.columns[1].width = Inches(3.0)

    # Right table — "INVESTMENT RETURNS" header
    _add_textbox(slide, Inches(7), Inches(1.15), Inches(5), Inches(0.3),
                 "INVESTMENT RETURNS", font_size=12, color=GOLD, bold=True)

    right_table = _add_table(
        slide, Inches(7), Inches(1.4), Inches(5.5), Inches(4.5),
        len(right_data), 2,
    )
    for i, (label, value) in enumerate(right_data):
        bg = DARK_800 if i % 2 == 0 else DARK_900
        _style_data_cell(right_table.cell(i, 0), label, color=LIGHT_GRAY,
                         alignment=PP_ALIGN.LEFT, bg=bg)
        _style_data_cell(right_table.cell(i, 1), value, color=WHITE,
                         bold=True, alignment=PP_ALIGN.RIGHT, bg=bg)

    right_table.columns[0].width = Inches(2.5)
    right_table.columns[1].width = Inches(3.0)

    # Footer
    _add_textbox(slide, Inches(0.8), Inches(7), Inches(11), Inches(0.3),
                 _footer_text(brand_config),
                 font_size=7, color=MEDIUM_GRAY, alignment=PP_ALIGN.CENTER)


def _build_returns_slide(prs, inputs, results, scenario_results=None, brand_config=None):
    """Slide 3: Investment Returns with optional scenario comparison."""
    if brand_config is None:
        brand_config = BrandConfig()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide)

    _add_textbox(slide, Inches(0.8), Inches(0.4), Inches(11), Inches(0.6),
                 "INVESTMENT RETURNS", font_size=28, color=WHITE,
                 font_name="Georgia", bold=True)

    shape = slide.shapes.add_shape(
        1, Inches(0.8), Inches(1.05), Inches(2.5), Inches(0.04)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = GOLD
    shape.line.fill.background()

    if scenario_results and "comparison" in scenario_results:
        _build_scenario_table(slide, scenario_results)
    else:
        _build_single_returns_table(slide, results)

    _add_textbox(slide, Inches(0.8), Inches(7), Inches(11), Inches(0.3),
                 _footer_text(brand_config),
                 font_size=7, color=MEDIUM_GRAY, alignment=PP_ALIGN.CENTER)


def _build_single_returns_table(slide, results):
    """Build returns table when no scenarios are available."""
    metrics = results.get("metrics", {})
    irr_d = metrics.get("irr", {})
    em_d = metrics.get("equity_multiple", {})
    dscr_d = metrics.get("dscr", {})

    data = [
        ["Metric", "Levered", "Unlevered"],
        ["IRR", _fmt_pct(irr_d.get("levered_irr")), _fmt_pct(irr_d.get("unlevered_irr"))],
        ["Equity Multiple", _fmt_multiple(em_d.get("levered_em")), _fmt_multiple(em_d.get("unlevered_em"))],
        ["Avg DSCR", f"{dscr_d.get('average_dscr', '—')}", "—"],
        ["Min DSCR", f"{dscr_d.get('minimum_dscr', '—')}", "—"],
    ]

    table = _add_table(slide, Inches(2), Inches(2), Inches(9), Inches(2.5),
                       len(data), 3)

    _style_header_row(table, 3)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Levered"
    table.cell(0, 2).text = "Unlevered"

    for r in range(1, len(data)):
        bg = DARK_800 if r % 2 == 0 else DARK_900
        _style_data_cell(table.cell(r, 0), data[r][0], color=GOLD,
                         bold=True, alignment=PP_ALIGN.LEFT, bg=bg)
        for c in range(1, 3):
            _style_data_cell(table.cell(r, c), data[r][c], color=WHITE,
                             bold=True, bg=bg)

    for i in range(3):
        table.columns[i].width = Inches(3)


def _build_scenario_table(slide, scenario_results):
    """Build Bull/Base/Bear comparison table."""
    comp = scenario_results.get("comparison", {})

    # Metric rows
    metric_rows = [
        ("Levered IRR", "levered_irr", _fmt_pct),
        ("Unlevered IRR", "unlevered_irr", _fmt_pct),
        ("Levered EM", "levered_em", _fmt_multiple),
        ("Unlevered EM", "unlevered_em", _fmt_multiple),
        ("NOI Year 1", "noi_year_1", _fmt_currency),
        ("NOI Exit Year", "noi_exit_year", _fmt_currency),
        ("Avg DSCR", "average_dscr", lambda v: f"{v:.2f}" if v else "—"),
        ("Going-In Cap", "going_in_cap", _fmt_pct),
    ]

    # Assumptions rows
    assumption_rows = [
        ("Rent Growth", "rent_growth", _fmt_pct),
        ("Exit Cap Rate", "exit_cap_rate", _fmt_pct),
        ("Avg Vacancy", "avg_vacancy", _fmt_pct),
    ]

    total_rows = 1 + len(metric_rows) + 1 + len(assumption_rows)  # header + metrics + divider + assumptions
    table = _add_table(slide, Inches(1.5), Inches(1.5), Inches(10), Inches(5),
                       total_rows, 4)

    # Header
    headers = ["", "Bull", "Base", "Bear"]
    for i, h in enumerate(headers):
        table.cell(0, i).text = h
    _style_header_row(table, 4)

    # Metric rows
    scenarios = ["bull", "base", "bear"]
    for r, (label, key, formatter) in enumerate(metric_rows):
        row_idx = r + 1
        bg = DARK_800 if r % 2 == 0 else DARK_900
        _style_data_cell(table.cell(row_idx, 0), label, color=GOLD,
                         bold=True, alignment=PP_ALIGN.LEFT, bg=bg)
        for c, scenario in enumerate(scenarios):
            val = comp.get(scenario, {}).get("metrics", {}).get(key)
            _style_data_cell(table.cell(row_idx, c + 1), formatter(val),
                             color=WHITE, bold=(scenario == "base"), bg=bg)

    # Divider row (assumptions label)
    div_idx = len(metric_rows) + 1
    _style_data_cell(table.cell(div_idx, 0), "KEY ASSUMPTIONS", color=GOLD,
                     bold=True, alignment=PP_ALIGN.LEFT, bg=DARK_700)
    for c in range(1, 4):
        _style_data_cell(table.cell(div_idx, c), "", bg=DARK_700)

    # Assumption rows
    for r, (label, key, formatter) in enumerate(assumption_rows):
        row_idx = div_idx + 1 + r
        bg = DARK_800 if r % 2 == 0 else DARK_900
        _style_data_cell(table.cell(row_idx, 0), label, color=LIGHT_GRAY,
                         alignment=PP_ALIGN.LEFT, bg=bg)
        for c, scenario in enumerate(scenarios):
            val = comp.get(scenario, {}).get("assumptions", {}).get(key)
            _style_data_cell(table.cell(row_idx, c + 1), formatter(val),
                             color=LIGHT_GRAY, bg=bg)

    # Column widths
    table.columns[0].width = Inches(2.5)
    for i in range(1, 4):
        table.columns[i].width = Inches(2.5)


def _build_cashflow_slide(prs, results, brand_config=None):
    """Slide 4: Annual cash flow bar chart."""
    if brand_config is None:
        brand_config = BrandConfig()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide)

    _add_textbox(slide, Inches(0.8), Inches(0.4), Inches(11), Inches(0.6),
                 "CASH FLOW PROJECTIONS", font_size=28, color=WHITE,
                 font_name="Georgia", bold=True)

    shape = slide.shapes.add_shape(
        1, Inches(0.8), Inches(1.05), Inches(3), Inches(0.04)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = GOLD
    shape.line.fill.background()

    # Build chart with matplotlib
    by_year = results.get("cashflow", {}).get("by_year", [])
    if not by_year:
        _add_textbox(slide, Inches(3), Inches(3.5), Inches(7), Inches(1),
                     "No annual cash flow data available.",
                     font_size=16, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)
        return

    years = [str(y.get("year", "")) for y in by_year]
    noi = [y.get("net_operating_income", 0) for y in by_year]
    lcf = [y.get("levered_cashflow", 0) for y in by_year]

    fig, ax = plt.subplots(figsize=(10, 4.5))
    fig.patch.set_facecolor(MPL_BG)
    ax.set_facecolor(MPL_BG)

    x_pos = range(len(years))
    bar_width = 0.35

    bars1 = ax.bar([x - bar_width / 2 for x in x_pos], noi, bar_width,
                   label="NOI", color=MPL_DARK, edgecolor=MPL_GOLD, linewidth=0.5)
    bars2 = ax.bar([x + bar_width / 2 for x in x_pos], lcf, bar_width,
                   label="Levered CF", color=MPL_GOLD)

    ax.set_xticks(list(x_pos))
    ax.set_xticklabels(years, color="white", fontsize=10)
    ax.yaxis.set_major_formatter(mticker.FuncFormatter(lambda x, _: f"${x:,.0f}"))
    ax.tick_params(colors="white")
    ax.spines["bottom"].set_color("#333333")
    ax.spines["left"].set_color("#333333")
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.legend(facecolor=MPL_BG, edgecolor="#333333", labelcolor="white",
              fontsize=9, loc="upper left")

    plt.tight_layout()

    # Save to buffer and add to slide
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=200, facecolor=MPL_BG,
                bbox_inches="tight", pad_inches=0.2)
    plt.close(fig)
    buf.seek(0)

    slide.shapes.add_picture(buf, Inches(1.5), Inches(1.5), Inches(10), Inches(5))

    _add_textbox(slide, Inches(0.8), Inches(7), Inches(11), Inches(0.3),
                 _footer_text(brand_config),
                 font_size=7, color=MEDIUM_GRAY, alignment=PP_ALIGN.CENTER)


def _build_risk_slide(prs, inputs, results, brand_config=None):
    """Slide 5: Key Assumptions & Risk Factors."""
    if brand_config is None:
        brand_config = BrandConfig()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide)

    _add_textbox(slide, Inches(0.8), Inches(0.4), Inches(11), Inches(0.6),
                 "KEY ASSUMPTIONS & RISKS", font_size=28, color=WHITE,
                 font_name="Georgia", bold=True)

    shape = slide.shapes.add_shape(
        1, Inches(0.8), Inches(1.05), Inches(3), Inches(0.04)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = GOLD
    shape.line.fill.background()

    ga = inputs.get("growth_assumptions", {})
    exit_a = inputs.get("exit_assumptions", {})
    debt = inputs.get("debt_terms", {})
    vac_curve = inputs.get("physical_vacancy_curve", [])
    avg_vac = (
        sum(e.get("vacancy_rate", 0) for e in vac_curve) / len(vac_curve)
        if vac_curve else 0.05
    )

    # Left column: Key Assumptions
    _add_textbox(slide, Inches(0.8), Inches(1.5), Inches(5), Inches(0.4),
                 "KEY ASSUMPTIONS", font_size=14, color=GOLD, bold=True)

    assumptions = [
        f"Rent Growth: {_fmt_pct(ga.get('annual_growth_rate'))}",
        f"Exit Cap Rate: {_fmt_pct(exit_a.get('exit_cap_rate'))}",
        f"Avg Physical Vacancy: {_fmt_pct(avg_vac)}",
        f"Loan Rate: {_fmt_pct(debt.get('rate'))}",
        f"Loan LTV: {_fmt_pct(debt.get('commitment', 0) / inputs.get('purchase_assumptions', {}).get('purchase_price', 1))}",
        f"Amortization: {debt.get('amort_years', '—')} years",
    ]

    # Add IO period if present
    io_months = debt.get("io_months", 0)
    if io_months:
        assumptions.append(f"Interest Only: {io_months} months")

    for i, text in enumerate(assumptions):
        _add_textbox(slide, Inches(1.2), Inches(2.1 + i * 0.45), Inches(5), Inches(0.4),
                     text, font_size=12, color=WHITE)

    # Right column: Risk Factors
    _add_textbox(slide, Inches(7), Inches(1.5), Inches(5), Inches(0.4),
                 "RISK CONSIDERATIONS", font_size=14, color=GOLD, bold=True)

    risks = [
        "Market rent growth may underperform assumptions",
        "Interest rate environment may impact refinancing",
        "Physical vacancy may exceed projections",
        "Operating expenses may increase above inflation",
        "Capital expenditure timing and costs may vary",
        "Exit cap rate expansion could reduce returns",
        "Regulatory changes may affect operating costs",
    ]

    for i, text in enumerate(risks):
        _add_textbox(slide, Inches(7.4), Inches(2.1 + i * 0.45), Inches(5), Inches(0.4),
                     f"  {text}", font_size=11, color=LIGHT_GRAY)
        # Gold bullet
        _add_textbox(slide, Inches(7), Inches(2.1 + i * 0.45), Inches(0.3), Inches(0.4),
                     "\u2022", font_size=11, color=GOLD)

    # Disclaimer
    _add_textbox(slide, Inches(0.8), Inches(6.2), Inches(11.5), Inches(0.8),
                 "This presentation is provided for informational purposes only and does not "
                 "constitute an offer to sell or a solicitation of an offer to buy any securities. "
                 "Past performance is not indicative of future results. All projections are estimates "
                 "based on assumptions that may not be realized.",
                 font_size=7, color=MEDIUM_GRAY)

    _add_textbox(slide, Inches(0.8), Inches(7), Inches(11), Inches(0.3),
                 _footer_text(brand_config),
                 font_size=7, color=MEDIUM_GRAY, alignment=PP_ALIGN.CENTER)


# ── Public API ────────────────────────────────────────────────────────────


def generate_presentation(
    results: Dict[str, Any],
    inputs: Dict[str, Any],
    output_path: str | Path,
    scenario_results: Optional[Dict[str, Any]] = None,
    brand_config: Optional[BrandConfig] = None,
) -> Path:
    """Generate a branded PowerPoint presentation.

    Args:
        results: Engine output dict
        inputs: Engine input dict
        output_path: Where to save the .pptx file
        scenario_results: Optional Bull/Base/Bear scenario comparison
        brand_config: Optional brand config for JV co-branding.
            If None, built from inputs (fund_assumptions.jv_partner_name).

    Returns:
        Path to the saved presentation
    """
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    if brand_config is None:
        brand_config = BrandConfig.from_inputs(inputs)

    prs = Presentation()
    # Set widescreen 16:9
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT

    _build_cover_slide(prs, inputs, results, brand_config)
    _build_summary_slide(prs, inputs, results, brand_config)
    _build_returns_slide(prs, inputs, results, scenario_results, brand_config)
    _build_cashflow_slide(prs, results, brand_config)
    _build_risk_slide(prs, inputs, results, brand_config)

    prs.save(str(output_path))
    return output_path
