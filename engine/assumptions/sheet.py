"""Assumptions sheet: export and import (ADR 0003 §8, slice 7).

Every run writes an assumptions sheet so an analyst can inspect each resolved
assumption and change it. The sheet is a flat view of the canonical
``assumptions.entries``: one row per entry with its value, provenance, the
resolver default, the analyst override and the delta from the default.

Two formats carry the same rows:

- ``assumptions.json``: ``write_assumptions_json`` / ``read_assumption_overrides``.
- An "Assumptions" Excel tab (an Excel table named ``tbl_assumptions``):
  ``add_assumptions_sheet`` / ``write_assumptions_xlsx`` /
  ``read_assumption_overrides_xlsx``.

The read side turns an edited sheet into the overrides ``resolve_assumptions``
takes: ``{key: {"value": ..., "note": ...}}`` sets an override and
``{key: None}`` clears one. The sheet carries each row's override as exported
(the read-only ``exported_override`` column), so the reader can tell what the
analyst changed:

- an Override cell changed from the export sets that override (it wins);
- otherwise an edited Value (different from the exported override, or from
  ``default_value`` when there was none) sets an override, so editing the Value
  column always works;
- a blanked Override cell whose row had an override clears it (``None``), as
  does a Value set back to ``default_value`` with the Override cell blank;
- an Override cell left as exported keeps that override.

Sheets written before the ``exported_override`` column (sheet version 1) are
read as before: a set Override wins, else a Value that differs from
``default_value`` becomes an override.

Rows are built generically from whatever an entry carries: scalar, vector and
dict values all work, and fields the sheet does not know about are kept in the
row's ``extra`` object. Nothing here special-cases an assumption key.
"""

from __future__ import annotations

import json
import math
from copy import deepcopy
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable, Mapping, Optional

SHEET_VERSION = "plat.assumptions.sheet/2"
SHEET_TITLE = "Assumptions"
TABLE_NAME = "tbl_assumptions"
DEFAULT_EDIT_NOTE = "edited in the assumptions sheet"

# Entry fields with a dedicated sheet column. Every other entry field is kept
# verbatim in the row's ``extra`` object. ``candidates`` stays in the
# canonical (it is the audit trail, not an editable value) and an entry's
# top-level ``cohort_overrides`` mirrors ``override.cohort_overrides``.
_ENTRY_COLUMNS = (
    "value",
    "unit",
    "source",
    "basis",
    "as_of",
    "confidence",
    "review_required",
    "default_value",
    "override",
)
_NOT_EXTRA = frozenset(_ENTRY_COLUMNS) | {"candidates", "cohort_overrides"}

# (row field, Excel header). The order is the column order of both formats.
COLUMNS: tuple[tuple[str, str], ...] = (
    ("key", "Key"),
    ("value", "Value"),
    ("unit", "Unit"),
    ("source", "Source"),
    ("basis", "Basis"),
    ("as_of", "As Of"),
    ("confidence", "Confidence"),
    ("review_required", "Review Required"),
    ("default_value", "Default Value"),
    ("delta", "Delta"),
    ("override_value", "Override"),
    ("override_note", "Note"),
    ("override_cohorts", "Override Cohorts"),
    ("guardrail_status", "Guardrail"),
    ("extra", "Extra"),
)
ROW_FIELDS = tuple(field for field, _ in COLUMNS)
# The override value each row carried when the sheet was written. It is
# read-only: the reader compares the Override cell with it to tell an edited, a
# cleared and an untouched override apart. It is a sheet column, not an entry
# field, so ``sheet_rows`` (one row per entry) does not carry it.
EXPORTED_OVERRIDE_FIELD = "exported_override"
SHEET_COLUMNS: tuple[tuple[str, str], ...] = COLUMNS + ((EXPORTED_OVERRIDE_FIELD, "Exported Override"),)
SHEET_FIELDS = tuple(field for field, _ in SHEET_COLUMNS)
_HEADER_TO_FIELD = {header.lower(): field for field, header in SHEET_COLUMNS}
_EDITABLE_FIELDS = ("override_value", "override_note")
# Columns whose cells may hold a number, a vector or an object.
_VALUE_FIELDS = frozenset(
    {"value", "default_value", "delta", "override_value", "override_cohorts", "extra", EXPORTED_OVERRIDE_FIELD}
)
_NUMBER_FORMAT_FIELDS = ("value", "default_value", "delta", "override_value", EXPORTED_OVERRIDE_FIELD)

_REL_TOL = 1e-9
_ABS_TOL = 1e-12


# ── value helpers ────────────────────────────────────────────────────────


def _number(value: Any) -> Optional[float]:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (int, float)):
        return float(value)
    return None


def _clean(number: float) -> float:
    """Round away float noise and normalize ``-0.0`` so deltas are stable."""
    rounded = round(number, 12)
    return 0.0 if rounded == 0 else rounded


def _same(a: Any, b: Any) -> bool:
    """Structural equality with a float tolerance (Excel round-trips floats)."""
    if a is None or b is None:
        return a is None and b is None
    if isinstance(a, (list, tuple)) or isinstance(b, (list, tuple)):
        if not isinstance(a, (list, tuple)) or not isinstance(b, (list, tuple)) or len(a) != len(b):
            return False
        return all(_same(x, y) for x, y in zip(a, b))
    if isinstance(a, Mapping) or isinstance(b, Mapping):
        if not isinstance(a, Mapping) or not isinstance(b, Mapping) or set(a) != set(b):
            return False
        return all(_same(a[k], b[k]) for k in a)
    na, nb = _number(a), _number(b)
    if na is not None and nb is not None:
        return math.isclose(na, nb, rel_tol=_REL_TOL, abs_tol=_ABS_TOL)
    return a == b


def delta(value: Any, default: Any) -> Any:
    """``value - default`` for numeric scalars, element-wise for numeric vectors.

    Anything else (dicts, text, mismatched shapes, missing values) has no
    delta and returns ``None``.
    """
    nv, nd = _number(value), _number(default)
    if nv is not None and nd is not None:
        return _clean(nv - nd)
    if isinstance(value, (list, tuple)) and isinstance(default, (list, tuple)) and len(value) == len(default):
        pairs = [(_number(v), _number(d)) for v, d in zip(value, default)]
        if pairs and all(v is not None and d is not None for v, d in pairs):
            return [_clean(v - d) for v, d in pairs]
    return None


def _jsonable(value: Any) -> Any:
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, Mapping):
        return {str(k): _jsonable(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [_jsonable(v) for v in value]
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _guardrail_status(guardrail: Any) -> Optional[str]:
    if guardrail is None:
        return None
    if isinstance(guardrail, Mapping):
        status = guardrail.get("status")
        return None if status is None else str(status)
    return str(guardrail)


# ── rows ─────────────────────────────────────────────────────────────────


def _entries(canonical: Mapping[str, Any]) -> Mapping[str, Any]:
    section = canonical.get("assumptions") if isinstance(canonical, Mapping) else None
    entries = section.get("entries") if isinstance(section, Mapping) else None
    return entries if isinstance(entries, Mapping) else {}


def has_assumptions(canonical: Any) -> bool:
    """True when the canonical carries an ``assumptions`` section."""
    return isinstance(canonical, Mapping) and isinstance(canonical.get("assumptions"), Mapping)


def _key_order() -> dict[str, int]:
    from engine.assumptions.resolver import ASSUMPTION_KEYS

    return {key: index for index, key in enumerate(ASSUMPTION_KEYS)}


def sorted_keys(keys: Iterable[str]) -> list[str]:
    """Registry order first (``ASSUMPTION_KEYS``), then any other key A→Z."""
    order = _key_order()
    return sorted(keys, key=lambda key: (order.get(key, len(order)), key))


def entry_row(key: str, entry: Mapping[str, Any]) -> dict[str, Any]:
    """One sheet row for one registry entry. Works for any entry shape."""
    entry = entry if isinstance(entry, Mapping) else {"value": entry}
    value = _jsonable(entry.get("value"))
    default_value = _jsonable(entry.get("default_value"))
    override = entry.get("override")
    override = override if isinstance(override, Mapping) else None
    extra = {
        str(name): _jsonable(entry[name])
        for name in sorted(entry, key=str)
        if name not in _NOT_EXTRA
    }
    row = {
        "key": key,
        "value": value,
        "unit": entry.get("unit"),
        "source": entry.get("source"),
        "basis": entry.get("basis"),
        "as_of": entry.get("as_of"),
        "confidence": entry.get("confidence"),
        "review_required": bool(entry.get("review_required", False)),
        "default_value": default_value,
        "delta": delta(value, default_value),
        "override_value": _jsonable(override.get("value")) if override else None,
        "override_note": override.get("note") if override else None,
        "override_cohorts": _jsonable(override.get("cohort_overrides")) if override else None,
        "guardrail_status": _guardrail_status(entry.get("guardrail")),
        "extra": extra or None,
    }
    return row


def sheet_rows(canonical: Mapping[str, Any]) -> list[dict[str, Any]]:
    """One row per assumption entry, in a stable key order."""
    entries = _entries(canonical)
    return [entry_row(key, entries[key]) for key in sorted_keys(entries)]


def document_rows(canonical: Mapping[str, Any]) -> list[dict[str, Any]]:
    """The rows a written sheet carries: ``sheet_rows`` plus ``exported_override``."""
    return [
        {**row, EXPORTED_OVERRIDE_FIELD: deepcopy(row["override_value"])}
        for row in sheet_rows(canonical)
    ]


SHEET_INSTRUCTIONS = (
    "Edit a row's value, or set its override_value (and override_note), then pass this file to "
    "--assumption-overrides or `runs/assumptions_sheet.py apply`. An override_value changed from "
    "exported_override wins over an edited value; with override_value left as exported, an edited "
    "value becomes the override. Blank override_value (null) to clear the row's override and "
    "return it to default_value. Do not edit exported_override."
)


def assumptions_document(canonical: Mapping[str, Any]) -> dict[str, Any]:
    """The ``assumptions.json`` payload: sheet rows plus section metadata."""
    section = canonical.get("assumptions") if has_assumptions(canonical) else {}
    metadata = canonical.get("metadata") if isinstance(canonical.get("metadata"), Mapping) else {}
    return {
        "sheet_version": SHEET_VERSION,
        "resolver_version": section.get("resolver_version"),
        "mode": section.get("mode"),
        "deal_id": metadata.get("deal_id"),
        "instructions": SHEET_INSTRUCTIONS,
        "columns": list(SHEET_FIELDS),
        "rows": document_rows(canonical),
    }


def write_assumptions_json(canonical: Mapping[str, Any], path: Path | str) -> Path:
    """Write ``assumptions.json`` for the canonical and return its path."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = json.dumps(assumptions_document(canonical), indent=2, ensure_ascii=False)
    path.write_text(payload + "\n", encoding="utf-8")
    return path


def sheet_path_for(output_path: Path | str) -> Path:
    """``deal.json`` → ``deal.assumptions.json`` in the same folder."""
    output_path = Path(output_path)
    return output_path.with_name(f"{output_path.stem}.assumptions.json")


# ── overrides from an edited sheet ───────────────────────────────────────


# What a row asks of the resolver: nothing, an override, or a clear.
_NO_CHANGE = object()
CLEAR = None  # ``{key: None}`` clears the key's override (resolve_assumptions contract)


def _row_override(row: Mapping[str, Any]) -> Any:
    """The override a row asks for: ``_NO_CHANGE``, ``CLEAR`` (None) or a dict.

    See the module docstring for the rules. A row without an
    ``exported_override`` field (a version-1 sheet or a hand-built row list) is
    read with the version-1 rules.
    """
    value = row.get("value")
    default_value = row.get("default_value")
    override_value = row.get("override_value")
    note = row.get("override_note")
    note = str(note).strip() if note is not None and str(note).strip() else None
    cohorts = row.get("override_cohorts")

    if EXPORTED_OVERRIDE_FIELD not in row:
        # Version 1: a set override wins, else a value that differs from the default.
        if override_value is not None:
            chosen, keep_cohorts = override_value, True
        elif value is not None and not _same(value, default_value):
            chosen, keep_cohorts = value, False
        else:
            return _NO_CHANGE
    else:
        exported = row.get(EXPORTED_OVERRIDE_FIELD)
        override_edited = override_value is not None and not _same(override_value, exported)
        baseline = exported if exported is not None else default_value
        value_edited = value is not None and not _same(value, baseline)
        if override_edited:
            chosen, keep_cohorts = override_value, True
        elif value_edited and override_value is None and exported is not None and _same(value, default_value):
            return CLEAR  # override blanked and the value set back to the default
        elif value_edited:
            chosen, keep_cohorts = value, True
        elif override_value is None and exported is not None:
            return CLEAR  # override blanked
        elif override_value is not None:
            chosen, keep_cohorts = override_value, True  # kept as exported
        else:
            return _NO_CHANGE
    out: dict[str, Any] = {"value": chosen, "note": note or DEFAULT_EDIT_NOTE}
    if keep_cohorts and cohorts:
        out["cohort_overrides"] = cohorts
    return out


def overrides_from_rows(rows: Iterable[Mapping[str, Any]]) -> dict[str, Optional[dict[str, Any]]]:
    """``{key: {"value", "note"[, "cohort_overrides"]}}`` for each edited row, and
    ``{key: None}`` for each row whose override was cleared."""
    out: dict[str, Optional[dict[str, Any]]] = {}
    seen: set[str] = set()
    for index, row in enumerate(rows):
        if not isinstance(row, Mapping):
            raise ValueError(f"assumptions sheet row {index} must be an object")
        key = row.get("key")
        if not isinstance(key, str) or not key.strip():
            raise ValueError(f"assumptions sheet row {index} has no key")
        key = key.strip()
        if key in seen:
            raise ValueError(f"assumptions sheet lists {key!r} twice")
        seen.add(key)
        override = _row_override(row)
        if override is not _NO_CHANGE:
            out[key] = override
    return {key: out[key] for key in sorted_keys(out)}


def overrides_from_document(data: Any) -> dict[str, Any]:
    """Overrides from an exported sheet document, a bare row list, or a plain
    ``{key: value | {"value": ..., "note": ...} | None}`` overrides object
    (returned as-is; ``None`` clears the key's override)."""
    if isinstance(data, list):
        return overrides_from_rows(data)
    if not isinstance(data, Mapping):
        raise ValueError("assumption overrides must be a JSON object or a list of sheet rows")
    if "rows" in data:
        rows = data["rows"]
        if not isinstance(rows, list):
            raise ValueError("assumptions sheet 'rows' must be a list")
        return overrides_from_rows(rows)
    return dict(data)


def read_assumption_overrides(path: Path | str) -> dict[str, Any]:
    """Overrides from an edited ``assumptions.json`` (or a plain overrides JSON).

    A row whose override was cleared comes back as ``{key: None}``.
    """
    with open(path, encoding="utf-8") as handle:
        data = json.load(handle)
    return overrides_from_document(data)


def load_assumption_overrides(path: Path | str) -> dict[str, Any]:
    """Overrides from a ``.json`` sheet/overrides file or an ``.xlsx``/``.xlsm`` workbook."""
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        return read_assumption_overrides_xlsx(path)
    return read_assumption_overrides(path)


def stamp_overrides(
    overrides: Optional[Mapping[str, Any]],
    canonical: Mapping[str, Any],
    *,
    edited_by: Optional[str] = None,
    edited_at: Optional[str] = None,
) -> dict[str, Any]:
    """Attach optional edit metadata (``edited_by`` / ``edited_at``) to overrides.

    ``overrides`` are about to be applied to ``canonical`` (an edited sheet,
    a plain overrides file):

    - an override equal to the one the canonical already records (same value,
      note and cohorts; for example an untouched row of a re-applied sheet)
      keeps that override's recorded metadata;
    - any other override (a new or changed edit) gets ``edited_by`` /
      ``edited_at`` when given, unless it carries its own;
    - a clear (``None``) is passed through.

    Nothing is added when there is no metadata to add, so overrides without
    metadata reach the resolver exactly as before. This never reads the clock:
    a caller that wants today's date passes it.
    """
    from engine.assumptions.registry import OVERRIDE_METADATA_FIELDS

    entries = _entries(canonical)
    given = {name: value for name, value in (("edited_by", edited_by), ("edited_at", edited_at)) if value}
    out: dict[str, Any] = {}
    for key, raw in (overrides or {}).items():
        if raw is None or (isinstance(raw, Mapping) and raw.get("value") is None):
            out[key] = raw
            continue
        override = dict(raw) if isinstance(raw, Mapping) else {"value": raw}
        entry = entries.get(key) if isinstance(entries.get(key), Mapping) else {}
        prior = entry.get("override") if isinstance(entry.get("override"), Mapping) else None
        unchanged = (
            prior is not None
            and _same(override.get("value"), prior.get("value"))
            and (override.get("note") is None or override.get("note") == prior.get("note"))
            and _same(override.get("cohort_overrides") or None, prior.get("cohort_overrides") or None)
        )
        metadata = (
            {name: prior[name] for name in OVERRIDE_METADATA_FIELDS if prior.get(name)}
            if unchanged
            else given
        )
        for name, value in metadata.items():
            override.setdefault(name, value)
        out[key] = override if (metadata or isinstance(raw, Mapping)) else raw
    return out


# ── Excel ────────────────────────────────────────────────────────────────


def _to_cell(value: Any) -> Any:
    """Scalars are written as themselves; vectors and objects as JSON text."""
    if isinstance(value, (list, tuple, Mapping)):
        return json.dumps(_jsonable(value), ensure_ascii=False)
    if isinstance(value, Decimal):
        return float(value)
    return value


def _from_cell(value: Any, *, field: str, key: Any) -> Any:
    """Decode a cell of a value column back to a number, vector or object."""
    if value is None:
        return None
    if not isinstance(value, str):
        return value
    text = value.strip()
    if not text:
        return None
    if text[0] in "[{":
        try:
            return json.loads(text)
        except json.JSONDecodeError as exc:
            raise ValueError(f"assumptions sheet {key!r} {field}: invalid JSON {text!r} ({exc})") from exc
    percent = text.endswith("%")
    try:
        number = float(text[:-1] if percent else text.replace(",", ""))
    except ValueError:
        return text
    return number / 100.0 if percent else number


def _number_format(unit: Any) -> Optional[str]:
    unit = str(unit or "")
    if unit.startswith("decimal"):
        return "0.00%"
    if unit.startswith("usd"):
        return '"$"#,##0.00'
    return None


# Register columns the Excel tab adds for reading (ADR 0003 §8, the
# assumption register). They are display-only: the reader maps headers to
# sheet fields and ignores these, so editing them changes nothing.
REGISTER_COLUMNS: tuple[tuple[str, str], ...] = (
    ("group", "Group"),
    ("default_source", "Default Source"),
    ("why", "Why This Default"),
    ("edited_by_at", "Edited By / At"),
    ("flags", "Flags"),
)
# The tab's column order: sheet fields (``SHEET_COLUMNS``) and register
# columns, arranged for reading left to right. Key stays in column A, and
# Note directly follows Override.
EXCEL_COLUMN_ORDER: tuple[str, ...] = (
    "key", "group", "value", "unit", "default_value", "default_source", "delta",
    "override_value", "override_note", "edited_by_at", "flags", "why",
    "source", "basis", "as_of", "confidence", "review_required",
    "override_cohorts", "guardrail_status", "extra", EXPORTED_OVERRIDE_FIELD,
)
_EXCEL_HEADERS = {**dict(SHEET_COLUMNS), **dict(REGISTER_COLUMNS)}
assert sorted(EXCEL_COLUMN_ORDER) == sorted(_EXCEL_HEADERS), "every sheet and register column is laid out once"

_HEX_EDITED_ROW = "FFF2CC"
_HEX_FLAG_REVIEW = "FBE3E0"
_HEX_FLAG_REVIEW_FONT = "9C2A1C"


def _register_cells(register_row: Mapping[str, Any]) -> dict[str, Any]:
    from engine.assumptions.register import flag_labels, guardrail_detail

    edit = register_row.get("user_edit") or {}
    who = " · ".join(str(part) for part in (edit.get("edited_by"), edit.get("edited_at")) if part)
    flags = "; ".join(flag_labels(register_row))
    detail = guardrail_detail(register_row)
    if detail:
        flags = f"{flags} ({detail})" if flags else detail
    return {
        "group": register_row.get("group"),
        "default_source": (register_row.get("default") or {}).get("source"),
        "why": (register_row.get("why") or {}).get("rationale"),
        "edited_by_at": who or None,
        "flags": flags or None,
    }


def add_assumptions_sheet(wb: Any, canonical: Mapping[str, Any], *, title: str = SHEET_TITLE) -> Any:
    """Add an "Assumptions" tab (Excel table ``tbl_assumptions``) to an openpyxl workbook.

    The tab is the assumption register in a workbook: rows sorted by group,
    then key order (``register.ordered_keys``), the sheet columns plus Group,
    Default Source, Why This Default, Edited By / At and Flags, the header
    frozen, rows the user edited highlighted and review flags marked. Override
    and Note (and Value) stay the editable cells; the reader
    (``read_assumption_overrides_xlsx``) is unchanged by the extra columns and
    the row order.
    """
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    from engine.assumptions.register import build_register
    from engine.brand import HEX_DARK_900, HEX_GOLD_LIGHT
    from engine.xlsx_helpers import set_col_widths, write_data_row, write_header_row

    if title in wb.sheetnames:
        del wb[title]
    ws = wb.create_sheet(title)
    register = build_register(canonical, include_candidates=False)
    by_key = {row["key"]: row for row in document_rows(canonical)}
    rows = [(by_key[reg["key"]], reg) for reg in register["rows"] if reg["key"] in by_key]
    column = {name: index + 1 for index, name in enumerate(EXCEL_COLUMN_ORDER)}

    ws["A1"] = "ASSUMPTIONS"
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color=HEX_DARK_900)
    ws["A2"] = (
        "Edit Value, or Override and Note, then pass this workbook to --assumption-overrides or "
        "`runs/assumptions_sheet.py apply`. A changed Override wins; with Override as exported, an "
        "edited Value becomes the override. Clear an Override cell to return the row to its default. "
        "Do not edit Exported Override. Vectors and objects are JSON text."
    )
    ws["A2"].font = Font(name="Calibri", size=10, italic=True, color=HEX_DARK_900)
    ws["A3"] = (
        "Rows are grouped by assumption group. Highlighted rows carry a user edit; a red Flags cell needs "
        "review. Group, Default Source, Why This Default, Edited By / At and Flags are for reading only."
    )
    ws["A3"].font = Font(name="Calibri", size=10, italic=True, color=HEX_DARK_900)

    header_row = 4
    write_header_row(ws, header_row, [_EXCEL_HEADERS[name] for name in EXCEL_COLUMN_ORDER])
    editable_fill = PatternFill(start_color=HEX_GOLD_LIGHT, end_color=HEX_GOLD_LIGHT, fill_type="solid")
    edited_fill = PatternFill(start_color=_HEX_EDITED_ROW, end_color=_HEX_EDITED_ROW, fill_type="solid")
    review_fill = PatternFill(start_color=_HEX_FLAG_REVIEW, end_color=_HEX_FLAG_REVIEW, fill_type="solid")
    editable_cols = [column[field] for field in _EDITABLE_FIELDS]
    wrap = Alignment(horizontal="left", vertical="center", wrap_text=True)

    for offset, (row, reg) in enumerate(rows, start=1):
        excel_row = header_row + offset
        cells = {**{field: _to_cell(row[field]) for field in SHEET_FIELDS}, **_register_cells(reg)}
        write_data_row(ws, excel_row, [cells[name] for name in EXCEL_COLUMN_ORDER], zebra=offset % 2 == 0)
        fmt = _number_format(row.get("unit"))
        for field in _NUMBER_FORMAT_FIELDS:
            cell = ws.cell(row=excel_row, column=column[field])
            if fmt and _number(cell.value) is not None:
                cell.number_format = fmt
        if reg.get("user_edit") is not None:
            for col in range(1, len(EXCEL_COLUMN_ORDER) + 1):
                ws.cell(row=excel_row, column=col).fill = edited_fill
            ws.cell(row=excel_row, column=column["key"]).font = Font(
                name="Calibri", size=10, bold=True, color=HEX_DARK_900
            )
        for col in editable_cols:
            ws.cell(row=excel_row, column=col).fill = editable_fill
        flagged = reg.get("review_required") or (reg.get("guardrail") or {}).get("state") in ("warned", "enforced")
        if flagged:
            for name in ("flags", "review_required"):
                cell = ws.cell(row=excel_row, column=column[name])
                cell.fill = review_fill
                cell.font = Font(name="Calibri", size=10, bold=True, color=_HEX_FLAG_REVIEW_FONT)
        for name in ("basis", "why", "flags"):
            ws.cell(row=excel_row, column=column[name]).alignment = wrap
    for col in editable_cols:
        ws.cell(row=header_row, column=col).font = Font(
            name="Calibri", size=10, bold=True, color=HEX_DARK_900
        )
        ws.cell(row=header_row, column=col).fill = editable_fill

    # An Excel table needs at least one data row; an empty sheet keeps a blank one.
    last_row = header_row + max(len(rows), 1)
    ref = f"A{header_row}:{get_column_letter(len(EXCEL_COLUMN_ORDER))}{last_row}"
    table = Table(displayName=TABLE_NAME, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
    ws.add_table(table)

    widths = {
        "key": 30, "group": 16, "value": 16, "unit": 26, "default_value": 16, "default_source": 24,
        "delta": 14, "override_value": 16, "override_note": 32, "edited_by_at": 22, "flags": 30,
        "why": 60, "source": 24, "basis": 70, "as_of": 12, "confidence": 12, "review_required": 10,
        "override_cohorts": 22, "guardrail_status": 14, "extra": 30, EXPORTED_OVERRIDE_FIELD: 18,
    }
    set_col_widths(ws, {column[name]: width for name, width in widths.items()})
    ws.freeze_panes = ws.cell(row=header_row + 1, column=2)
    return ws


def write_assumptions_xlsx(canonical: Mapping[str, Any], path: Path | str) -> Path:
    """Write a workbook holding only the "Assumptions" tab."""
    from openpyxl import Workbook

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    default = wb.active
    add_assumptions_sheet(wb, canonical)
    wb.remove(default)
    wb.save(path)
    return path


def _rows_from_table(headers: list[Any], cells: Iterable[list[Any]]) -> list[dict[str, Any]]:
    fields: list[Optional[str]] = [
        _HEADER_TO_FIELD.get(str(header).strip().lower()) if header is not None else None
        for header in headers
    ]
    if "key" not in fields:
        raise ValueError("assumptions sheet has no Key column")
    rows: list[dict[str, Any]] = []
    for raw in cells:
        row: dict[str, Any] = {}
        for field, cell in zip(fields, raw):
            if field is not None:
                row[field] = cell
        key = row.get("key")
        if key is None or (isinstance(key, str) and not key.strip()):
            continue
        for field in _VALUE_FIELDS:
            if field in row:
                row[field] = _from_cell(row[field], field=field, key=key)
        rows.append(row)
    # The tab may show rows in any order (grouped by the register); rows read
    # back follow the sheet's key order, like ``document_rows``.
    order = {key: index for index, key in enumerate(sorted_keys({str(row["key"]) for row in rows}))}
    return sorted(rows, key=lambda row: order[str(row["key"])])


def read_assumption_rows_xlsx(path: Path | str) -> list[dict[str, Any]]:
    """Sheet rows from a workbook's ``tbl_assumptions`` table.

    Reads through ``engine.excel_io`` like every other workbook table. When the
    table is gone (some editors drop Excel tables), the "Assumptions" tab is
    scanned for its header row instead.
    """
    from engine.excel_io import _is_blank_row, _normalize_cell, _read_excel_table

    path = Path(path)
    try:
        table = _read_excel_table(path, TABLE_NAME)
    except KeyError:
        table = None
    if table is not None:
        return _rows_from_table(list(table.headers), table.rows)

    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    try:
        if SHEET_TITLE not in wb.sheetnames:
            raise KeyError(f"workbook has no {TABLE_NAME} table or {SHEET_TITLE!r} tab: {path}")
        values = [
            [_normalize_cell(v) for v in row]
            for row in wb[SHEET_TITLE].iter_rows(values_only=True)
        ]
    finally:
        wb.close()
    for index, row in enumerate(values):
        if row and isinstance(row[0], str) and row[0].strip().lower() == "key":
            body = [r for r in values[index + 1:] if not _is_blank_row(r)]
            return _rows_from_table(row, body)
    raise KeyError(f"{SHEET_TITLE!r} tab has no Key header row: {path}")


def read_assumption_overrides_xlsx(path: Path | str) -> dict[str, Optional[dict[str, Any]]]:
    """Overrides from an edited "Assumptions" tab (``{key: None}`` for a cleared override)."""
    return overrides_from_rows(read_assumption_rows_xlsx(path))
