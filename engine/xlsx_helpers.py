"""
Shared Excel (openpyxl) Helper Functions
==========================================
Common utilities for writing styled rows and setting column widths
in openpyxl workbooks. Used by portfolio_xlsx and am_overlay_xlsx.
"""
from __future__ import annotations

from typing import Dict

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from engine.brand import HEX_DARK_900, HEX_LIGHT_GRAY, HEX_WHITE


# Pre-built style objects
_HEADER_FILL = PatternFill(start_color=HEX_DARK_900, end_color=HEX_DARK_900, fill_type="solid")
_LIGHT_FILL = PatternFill(start_color=HEX_LIGHT_GRAY, end_color=HEX_LIGHT_GRAY, fill_type="solid")
_WHITE_FILL = PatternFill(start_color=HEX_WHITE, end_color=HEX_WHITE, fill_type="solid")
_VALUE_FONT = Font(name="Calibri", size=10, color=HEX_DARK_900)

RIGHT = Alignment(horizontal="right", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")
CENTER = Alignment(horizontal="center", vertical="center")


def set_col_widths(ws, widths: Dict[int, float]):
    """Set column widths on a worksheet.

    Args:
        ws: openpyxl worksheet
        widths: mapping of 1-based column number → width
    """
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w


def write_header_row(ws, row: int, values: list, start_col: int = 1):
    """Write a dark header row with white bold text.

    Args:
        ws: openpyxl worksheet
        row: 1-based row number
        values: list of header strings
        start_col: 1-based starting column (default 1)
    """
    for i, val in enumerate(values):
        cell = ws.cell(row=row, column=start_col + i, value=val)
        cell.font = Font(name="Calibri", size=10, bold=True, color=HEX_WHITE)
        cell.fill = _HEADER_FILL
        cell.alignment = CENTER


def write_data_row(ws, row: int, values: list, start_col: int = 1, zebra: bool = False):
    """Write a data row with optional zebra striping.

    Args:
        ws: openpyxl worksheet
        row: 1-based row number
        values: list of cell values
        start_col: 1-based starting column (default 1)
        zebra: if True, use light gray background
    """
    fill = _LIGHT_FILL if zebra else _WHITE_FILL
    for i, val in enumerate(values):
        cell = ws.cell(row=row, column=start_col + i, value=val)
        cell.font = _VALUE_FONT
        cell.fill = fill
        cell.alignment = RIGHT if isinstance(val, (int, float)) else LEFT
